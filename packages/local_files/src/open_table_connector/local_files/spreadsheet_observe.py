"""Independent, bounded physical observations for XLSX layout state."""

from __future__ import annotations

import io
import re
import xml.etree.ElementTree as ET
from collections.abc import Mapping
from typing import Any
from zipfile import BadZipFile, ZipFile

from open_table_connector.contract import PROVIDER_EXCEL, SCHEME_XLSX
from open_table_connector.spreadsheets import ArtifactLimits
from open_table_connector.spreadsheets.observations import decode_observation

from .spreadsheet_verify import S, _coord, _decimal

_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_CELL = re.compile(r"^([A-Z]+)([1-9][0-9]*)$")
_ALL_STYLE_FIELDS = (
    "bold", "italic", "font_size", "foreground", "fill", "number_format",
    "horizontal", "vertical", "wrap_text", "shrink_to_fit", "text_layout", "border",
)


def _limits(value: ArtifactLimits | Mapping[str, Any] | None) -> dict[str, int]:
    if value is None:
        value = ArtifactLimits()
    result = {name: getattr(value, name) for name in ArtifactLimits.__dataclass_fields__} if isinstance(value, ArtifactLimits) else dict(value)
    allowed = set(ArtifactLimits.__dataclass_fields__)
    if set(result) - allowed or any(type(item) is not int or item <= 0 for item in result.values()):
        raise ValueError("invalid XLSX observation limits")
    return {name: result.get(name, getattr(ArtifactLimits(), name)) for name in allowed}


def _zip_parts(data: bytes, limits: Mapping[str, int]) -> dict[str, bytes]:
    if not isinstance(data, bytes) or len(data) > limits["archive_bytes"]:
        raise ValueError("XLSX snapshot exceeds archive limit")
    parts: dict[str, bytes] = {}
    total = 0
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            if len(archive.infolist()) > limits["zip_members"]:
                raise ValueError("XLSX snapshot exceeds member limit")
            for info in archive.infolist():
                if info.filename in parts or info.flag_bits & 1 or info.file_size > limits["member_bytes"]:
                    raise ValueError("unsafe XLSX archive member")
                total += info.file_size
                if total > limits["decompressed_bytes"]:
                    raise ValueError("XLSX snapshot exceeds decompressed limit")
                parts[info.filename] = archive.read(info)
    except (BadZipFile, OSError) as exc:
        raise ValueError("invalid XLSX archive") from exc
    return parts


def _xml(parts: Mapping[str, bytes], name: str) -> ET.Element:
    try:
        return ET.fromstring(parts[name])
    except (KeyError, ET.ParseError) as exc:
        raise ValueError(f"invalid or missing XLSX XML part: {name}") from exc


def _sheet_part(parts: Mapping[str, bytes], worksheet_id: str) -> tuple[str, ET.Element]:
    workbook = _xml(parts, "xl/workbook.xml")
    rels = _xml(parts, "xl/_rels/workbook.xml.rels")
    relationship = {
        item.attrib["Id"]: item.attrib["Target"]
        for item in rels.findall(f"{_PKG_REL}Relationship")
    }
    for sheet in workbook.findall(f"{_MAIN}sheets/{_MAIN}sheet"):
        if sheet.attrib.get("sheetId") == worksheet_id or sheet.attrib.get("name") == worksheet_id:
            rid = sheet.attrib.get(f"{_REL}id")
            target = relationship.get(rid or "")
            if target is None:
                raise ValueError("worksheet relationship is missing")
            part = target.lstrip("/")
            if not part.startswith("xl/"):
                part = "xl/" + part
            return part, sheet
    raise ValueError("worksheet identity was not found")


def _bounds(address: str) -> tuple[int, int, int, int]:
    match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)(?::([A-Z]+)([1-9][0-9]*))?", address.upper())
    if match is None:
        raise ValueError("invalid A1 range")
    start = _coord(match.group(1) + match.group(2))
    end = _coord((match.group(3) or match.group(1)) + (match.group(4) or match.group(2)))
    if end[0] < start[0] or end[1] < start[1]:
        raise ValueError("range end precedes start")
    return start[0], start[1], end[0], end[1]


def _raw_cells(sheet_xml: ET.Element) -> dict[str, ET.Element]:
    result: dict[str, ET.Element] = {}
    for cell in sheet_xml.findall(f".//{_MAIN}c"):
        ref = cell.attrib.get("r")
        if not ref or ref in result:
            raise ValueError("duplicate or missing cell coordinate")
        result[ref] = cell
    return result


def _source(sources: dict[str, Any], *, ref: str, kind: str, raw: str | None = None) -> str:
    sources[ref] = {"kind": kind, **({"raw_ref": raw} if raw else {})}
    return ref


def _effective_mode(wrap: bool, shrink: bool) -> str:
    if wrap and shrink:
        return "native_combination"
    if wrap:
        return "wrap"
    if shrink:
        return "shrink_to_fit"
    return "no_wrap"


def _observe_color(color: Any) -> Any:
    if color is None:
        return None
    color_type = getattr(color, "type", None)
    if color_type == "rgb" and color.rgb:
        value = str(color.rgb).upper()
        if len(value) == 8 and value[:2] == "FF":
            value = value[2:]
        return {"type": "rgb", "value": "#" + value}
    if color_type == "theme":
        return {"type": "theme", "value": color.theme, "tint": color.tint}
    if color_type == "indexed":
        return {"type": "indexed", "value": color.indexed, "tint": color.tint}
    if color_type == "auto":
        return {"type": "auto", "value": True}
    return None


def _cell_fields(cell: Any, raw: ET.Element | None, fields: list[str], sources: dict[str, Any], address: str, style_alignments: Mapping[str, ET.Element]) -> dict[str, Any]:
    raw_style = raw.attrib.get("s") if raw is not None else None
    style_kind = "explicit" if raw_style is not None else "default"
    raw_alignment = style_alignments.get(raw_style or "")
    result: dict[str, Any] = {}
    for field in fields:
        source_kind = style_kind
        if field in {"wrap_text", "shrink_to_fit", "horizontal", "vertical"}:
            source_kind = "explicit" if raw_alignment is not None and ("wrapText" in raw_alignment.attrib or "shrinkToFit" in raw_alignment.attrib or "horizontal" in raw_alignment.attrib or "vertical" in raw_alignment.attrib) else "default"
        source_ref = _source(sources, ref=f"cell:{address}:{field}", kind=source_kind, raw=f"xl/worksheets/{address}")
        if field == "bold":
            effective = bool(cell.font.bold)
        elif field == "italic":
            effective = bool(cell.font.italic)
        elif field == "font_size":
            effective = _decimal(cell.font.sz) if cell.font.sz else None
        elif field == "foreground":
            effective = _observe_color(cell.font.color)
        elif field == "fill":
            effective = _observe_color(cell.fill.fgColor) if cell.fill.patternType == "solid" else None
        elif field == "number_format":
            effective = {"storage_kind": "builtin" if cell.number_format in {"General", "0", "0.00", "0%", "m/d/yy"} else "custom", "code": cell.number_format, "builtin_id": None, "locale": None, "date_system": None}
        elif field == "horizontal":
            effective = cell.alignment.horizontal
        elif field == "vertical":
            effective = cell.alignment.vertical
        elif field == "wrap_text":
            effective = bool(cell.alignment.wrap_text)
        elif field == "shrink_to_fit":
            effective = bool(cell.alignment.shrink_to_fit)
        elif field == "text_layout":
            effective = _effective_mode(bool(cell.alignment.wrap_text), bool(cell.alignment.shrink_to_fit))
        elif field == "border":
            effective = {
                edge: {"style": getattr(cell.border, edge).style or "none", "color": _observe_color(getattr(cell.border, edge).color)}
                for edge in ("top", "bottom", "left", "right")
            }
        else:
            raise ValueError(f"unsupported style field: {field}")
        result[field] = {"effective": effective, "source_ref": source_ref}
    return result


def observe_xlsx(*, data: bytes, target: Mapping[str, Any], selector: Mapping[str, Any], limits: ArtifactLimits | Mapping[str, Any] | None = None) -> dict[str, Any]:
    """Read persisted XLSX bytes and return a validated physical observation."""

    if not isinstance(target, Mapping):
        raise ValueError("observation target must be an object")
    target = dict(target)
    target.setdefault("provider", PROVIDER_EXCEL)
    target.setdefault("resource", target.get("uri", SCHEME_XLSX))
    target.setdefault("worksheet_id", str(selector.get("worksheet_id", selector.get("sheet", "1"))))
    if not isinstance(selector, Mapping):
        raise ValueError("observation selector must be an object")
    bounds = _limits(limits)
    parts = _zip_parts(data, bounds)
    sheet_part, sheet_meta = _sheet_part(parts, str(target["worksheet_id"]))
    sheet_xml = _xml(parts, sheet_part)
    raw_cells = _raw_cells(sheet_xml)
    styles_root = _xml(parts, "xl/styles.xml")
    style_alignments = {}
    cell_xfs = styles_root.find(f"{_MAIN}cellXfs")
    if cell_xfs is not None:
        for index, xf in enumerate(cell_xfs.findall(f"{_MAIN}xf")):
            for alignment in xf.findall(f"{_MAIN}alignment"):
                if any(name in alignment.attrib for name in ("wrapText", "shrinkToFit", "horizontal", "vertical")):
                    style_alignments[str(index)] = alignment
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    try:
        sheet_name = sheet_meta.attrib.get("name")
        if sheet_name not in book.sheetnames:
            raise ValueError("worksheet title is missing")
        sheet = book[sheet_name]
        operation = selector.get("operation", "range.style.read")
        sources: dict[str, Any] = {}
        if operation == "range.style.read":
            address = selector.get("address")
            if not isinstance(address, str):
                raise ValueError("style observation requires address")
            start_row, start_col, end_row, end_col = _bounds(address)
            fields = selector.get("fields")
            fields = list(_ALL_STYLE_FIELDS if fields is None else fields)
            if not fields or len(set(fields)) != len(fields) or any(field not in _ALL_STYLE_FIELDS for field in fields):
                raise ValueError("invalid style observation fields")
            if (end_row - start_row + 1) * (end_col - start_col + 1) > bounds["cells"]:
                raise ValueError("observation exceeds cell limit")
            cells = []
            for row in range(start_row, end_row + 1):
                for column in range(start_col, end_col + 1):
                    coordinate = sheet.cell(row, column).coordinate
                    cells.append({"row": row, "column": column, "fields": _cell_fields(sheet.cell(row, column), raw_cells.get(coordinate), fields, sources, coordinate, style_alignments)})
            payload = {"cells": cells, "sources": sources}
            coverage = {"complete": True, "range": address.upper(), "fields": fields}
            kind = "spreadsheet.range.style.observation/1.0"
        elif operation == "worksheet.config.read":
            rows = selector.get("rows")
            columns = selector.get("columns")
            if not isinstance(rows, (list, tuple)) or not rows or not isinstance(columns, (list, tuple)) or not columns:
                raise ValueError("config observation requires rows and columns")
            if any(type(row) is not int or row < 1 for row in rows):
                raise ValueError("invalid config rows")
            row_values = {str(row): {"height": sheet.row_dimensions[row].height, "hidden": bool(sheet.row_dimensions[row].hidden), "source_ref": _source(sources, ref=f"row:{row}", kind="explicit" if sheet.row_dimensions[row].height is not None or sheet.row_dimensions[row].hidden else "default")} for row in rows}
            col_values = {}
            for column in columns:
                if not isinstance(column, str) or not re.fullmatch(r"[A-Z]{1,3}", column.upper()):
                    raise ValueError("invalid config column")
                dimension = sheet.column_dimensions[column.upper()]
                col_values[column.upper()] = {"width": dimension.width, "hidden": bool(dimension.hidden), "best_fit": bool(dimension.bestFit), "source_ref": _source(sources, ref=f"column:{column.upper()}", kind="explicit" if dimension.width is not None or dimension.hidden or dimension.bestFit else "default")}
            fields = list(selector.get("view_fields") or ("show_gridlines", "zoom_percent", "frozen_rows", "frozen_columns"))
            view = {"show_gridlines": sheet.sheet_view.showGridLines if sheet.sheet_view.showGridLines is not None else True, "zoom_percent": sheet.sheet_view.zoomScale, "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None}
            payload = {"rows": row_values, "columns": col_values, "view": view, "sources": sources}
            coverage = {"complete": True, "rows": list(rows), "columns": [str(item).upper() for item in columns], "fields": ["row_heights", "column_sizes", "gridlines", *fields]}
            kind = "spreadsheet.worksheet.config.observation/1.0"
        else:
            raise ValueError(f"unsupported XLSX observation operation: {operation}")
        observation = {"kind": kind, "target": target, "coverage": coverage, "physical": payload, "provenance": {"revision": "sha256:" + __import__("hashlib").sha256(data).hexdigest(), "source": "xlsx-xml"}}
        return decode_observation(observation)
    finally:
        book.close()


__all__ = ["observe_xlsx"]
