import json

from open_table_connector.cli.__main__ import main


def test_spreadsheet_batch_dry_run_then_commit_and_read(tmp_path, capsys):
    destination = tmp_path / "report.xlsx"
    commands = tmp_path / "commands.json"
    commands.write_text(
        json.dumps(
            {
                "version": "1.0",
                "changes": [
                    {"operation_id": "worksheet.create", "target_key": "Report", "arguments": {}},
                    {
                        "operation_id": "range.write",
                        "target_key": "Report",
                        "arguments": {"address": "A1", "values": [["value"]]},
                    },
                ],
            }
        )
    )
    argv = [
        "spreadsheet",
        "batch",
        "--uri",
        destination.as_uri(),
        "--create",
        "--commands",
        str(commands),
    ]
    assert main([*argv, "--dry-run"]) == 0
    assert json.loads(capsys.readouterr().out)["outcome"] == "planned"
    assert not destination.exists()
    assert main(argv) == 0
    assert json.loads(capsys.readouterr().out)["commit"] == "committed"
    assert (
        main(
            [
                "spreadsheet",
                "read",
                "--uri",
                destination.as_uri(),
                "--sheet",
                "Report",
                "--range",
                "A1",
            ]
        )
        == 0
    )
    assert json.loads(capsys.readouterr().out)["value"] == [["value"]]


def test_batch_rejects_unknown_keys_before_creating_file(tmp_path, capsys):
    destination = tmp_path / "report.xlsx"
    commands = tmp_path / "bad.json"
    commands.write_text(json.dumps({"version": "1.0", "changes": [], "ignored": True}))
    assert (
        main(
            [
                "spreadsheet",
                "batch",
                "--uri",
                destination.as_uri(),
                "--create",
                "--commands",
                str(commands),
            ]
        )
        != 0
    )
    assert not destination.exists()


def test_standalone_create_commits_and_verify_uses_retained_intent(tmp_path, capsys):
    destination = tmp_path / "report.xlsx"
    assert (
        main(
            [
                "spreadsheet",
                "operation",
                "--uri",
                destination.as_uri(),
                "--create",
                "--operation",
                "worksheet.create",
                "--sheet",
                "Report",
            ]
        )
        == 0
    )
    payload = json.loads(capsys.readouterr().out)
    assert payload["commit"] == "committed"
    expected = tmp_path / "expected.json"
    expected.write_text(json.dumps(payload["receipts"][0]["details"]["expected"]))
    assert (
        main(
            [
                "spreadsheet",
                "verify",
                "--uri",
                destination.as_uri(),
                "--profile",
                "literal-artifact/1.0",
                "--expected",
                str(expected),
            ]
        )
        == 0
    )
    assert json.loads(capsys.readouterr().out)["verification"] == "passed"
    assert (
        main(
            [
                "spreadsheet",
                "verify",
                "--uri",
                destination.as_uri(),
                "--profile",
                "literal-artifact/1.0",
            ]
        )
        == 5
    )
    assert json.loads(capsys.readouterr().err)["error"]["code"]


def _create_report(destination, capsys):
    assert main(
        [
            "spreadsheet",
            "operation",
            "--uri",
            destination.as_uri(),
            "--create",
            "--operation",
            "worksheet.create",
            "--sheet",
            "Report",
        ]
    ) == 0
    capsys.readouterr()


def test_style_shortcut_commits_and_persists(tmp_path, capsys):
    destination = tmp_path / "style.xlsx"
    _create_report(destination, capsys)
    assert main(["spreadsheet", "write", "--uri", destination.as_uri(), "--sheet", "Report", "--range", "A1", "--values", '[["value"]]']) == 0
    assert json.loads(capsys.readouterr().out)["commit"] == "committed"
    assert main(["spreadsheet", "style", "--uri", destination.as_uri(), "--sheet", "Report", "--range", "A1", "--bold"]) == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["commit"] == "committed"
    assert payload["verification"] == "passed"


def test_format_shortcut_commits_and_persists(tmp_path, capsys):
    destination = tmp_path / "format.xlsx"
    _create_report(destination, capsys)
    assert main(["spreadsheet", "format", "--uri", destination.as_uri(), "--sheet", "Report", "--range", "A1", "--pattern", "#,##0.00"]) == 0
    assert json.loads(capsys.readouterr().out)["commit"] == "committed"


def test_worksheet_shortcut_commits_and_persists(tmp_path, capsys):
    destination = tmp_path / "worksheet.xlsx"
    assert main(["spreadsheet", "worksheet", "create", "--uri", destination.as_uri(), "--create", "--sheet", "Report", "--name", "Summary"]) == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["commit"] == "committed"
    from openpyxl import load_workbook

    workbook = load_workbook(destination, read_only=True)
    assert "Summary" in workbook.sheetnames
    workbook.close()


def test_shortcut_dry_run_does_not_mutate(tmp_path, capsys):
    destination = tmp_path / "dry-run.xlsx"
    _create_report(destination, capsys)
    before = destination.read_bytes()
    assert main(["spreadsheet", "style", "--uri", destination.as_uri(), "--sheet", "Report", "--range", "A1", "--bold", "--dry-run"]) == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["outcome"] == "planned"
    assert destination.read_bytes() == before
