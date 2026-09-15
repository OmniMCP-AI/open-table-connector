"""Independent, bounded verification of the literal XLSX physical profile.

No writer/session state is consulted: the caller supplies a captured expectation.
"""

from __future__ import annotations

import hashlib
import io
import json
import math
import posixpath
import re
import struct
import xml.etree.ElementTree as ET
import zlib
from collections.abc import Iterator, Mapping
from decimal import Decimal
from typing import Any, NoReturn
from zipfile import BadZipFile, ZipFile, ZipInfo

from open_table_connector.contract import ConnectorError, ConnectorErrorCode

S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P = "http://schemas.openxmlformats.org/package/2006/relationships"
D = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
C = "http://schemas.openxmlformats.org/package/2006/content-types"
_DEFAULTS = dict(
    sheets=128,
    cells=250000,
    text_bytes=64 * 1024**2,
    images=128,
    image_bytes=16 * 1024**2,
    total_image_bytes=128 * 1024**2,
    image_pixels=40000000,
    archive_bytes=256 * 1024**2,
    zip_members=10000,
    member_bytes=128 * 1024**2,
    decompressed_bytes=512 * 1024**2,
)


def _fail(
    reason: str, code: ConnectorErrorCode = ConnectorErrorCode.READBACK_MISMATCH, **details: Any
) -> NoReturn:
    raise ConnectorError(
        code, "XLSX physical verification failed", {"reason": "artifact." + reason, **details}
    )


def _decimal(value: Any) -> float:
    number = float(value)
    if not math.isfinite(number):
        _fail("nonfinite_dimension")
    return float(Decimal(str(value)).quantize(Decimal("0.000001")))


def _color(color: Any) -> str | None:
    if color is None:
        return None
    if color.type != "rgb":
        _fail("unsupported_color")
    value = color.rgb.upper()
    if len(value) == 8:
        if value[:2] != "FF":
            _fail("unsupported_opacity")
        value = value[2:]
    return "#" + value


def _cell_style(cell: Any) -> dict[str, Any]:
    edges = {}
    for name in ("left", "right", "top", "bottom"):
        edge = getattr(cell.border, name)
        if edge is not None and edge.style:
            if edge.style not in ("thin", "medium", "thick", "dashed", "dotted", "double"):
                _fail("unsupported_border")
            edges[name] = {"style": edge.style, "color": _color(edge.color)}
    if cell.border.diagonalUp or cell.border.diagonalDown:
        _fail("unsupported_border")
    font_color = cell.font.color
    # The openpyxl default font is theme-based and is not a declared sRGB override.
    foreground = _color(font_color) if font_color is not None and font_color.type == "rgb" else None
    return dict(
        family=cell.font.name,
        size=_decimal(cell.font.sz) if cell.font.sz else None,
        bold=bool(cell.font.b),
        italic=bool(cell.font.i),
        foreground=foreground,
        fill=_color(cell.fill.fgColor) if cell.fill.patternType == "solid" else None,
        vertical=cell.alignment.vertical,
        horizontal=cell.alignment.horizontal,
        wrap_text=bool(cell.alignment.wrap_text),
        border=edges,
    )


def _sheet_config(ws: Any) -> dict[str, Any]:
    from openpyxl.worksheet.print_settings import PrintArea

    raw_area = ws.print_area
    if isinstance(raw_area, str):
        area = list(PrintArea.from_string(raw_area).ranges) if raw_area else []
    else:
        area = list(raw_area.ranges) if hasattr(raw_area, "ranges") else list(raw_area or [])
    return dict(
        visibility=ws.sheet_state,
        row_heights={
            str(k): _decimal(v.height) for k, v in ws.row_dimensions.items() if v.height is not None
        },
        column_widths={k: _decimal(v.width) for k, v in ws.column_dimensions.items()},
        show_gridlines=ws.sheet_view.showGridLines
        if ws.sheet_view.showGridLines is not None
        else True,
        freeze_panes=str(ws.freeze_panes) if ws.freeze_panes else None,
        print_area=sorted(str(v).replace("$", "") for v in area),
        orientation=ws.page_setup.orientation,
        paper_size=ws.page_setup.paperSize,
        fit_to_page=bool(
            ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage
        ),
        fit_width=ws.page_setup.fitToWidth,
        fit_height=ws.page_setup.fitToHeight,
        horizontal_centered=bool(ws.print_options.horizontalCentered),
        margins={
            k: _decimal(getattr(ws.page_margins, k))
            for k in ("left", "right", "top", "bottom", "header", "footer")
        },
    )


def _coord(value: str) -> tuple[int, int]:
    match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", value)
    if not match:
        _fail("invalid_coordinate")
    assert match is not None
    col = 0
    for char in match[1]:
        col = col * 26 + ord(char) - 64
    row = int(match[2])
    if col > 16384 or row > 1048576:
        _fail("invalid_coordinate")
    return row, col


def _part_kind(name: str) -> str:
    fixed = {
        "[Content_Types].xml": "types",
        "_rels/.rels": "rels",
        "docProps/core.xml": "core",
        "docProps/app.xml": "app",
        "xl/workbook.xml": "workbook",
        "xl/styles.xml": "styles",
        "xl/sharedStrings.xml": "sharedStrings",
        "xl/_rels/workbook.xml.rels": "rels",
    }
    if name in fixed:
        return fixed[name]
    for pattern, kind in [
        (r"xl/worksheets/sheet[0-9]+\.xml", "worksheet"),
        (r"xl/worksheets/_rels/sheet[0-9]+\.xml\.rels", "rels"),
        (r"xl/theme/theme[0-9]+\.xml", "theme"),
        (r"xl/drawings/drawing[0-9]+\.xml", "drawing"),
        (r"xl/drawings/_rels/drawing[0-9]+\.xml\.rels", "rels"),
        (r"xl/media/image[0-9]+\.(png|jpeg|jpg)", "image"),
    ]:
        if re.fullmatch(pattern, name):
            return kind
    _fail("unsupported_part")
    raise AssertionError


def _verify_snapshot(
    data: bytes, expected: Mapping[str, Any], limits: Any = None
) -> Mapping[str, Any]:
    """Return complete physical evidence, or raise a bounded ConnectorError."""
    bounds = dict(_DEFAULTS)
    if limits is not None:
        supplied = (
            dict(limits)
            if isinstance(limits, Mapping)
            else {k: getattr(limits, k) for k in bounds if hasattr(limits, k)}
        )
        if set(supplied) - set(bounds):
            _fail("invalid_limits", ConnectorErrorCode.CONFIGURATION)
        bounds.update(supplied)
    if any(type(v) is not int or v <= 0 for v in bounds.values()):
        _fail("invalid_limits", ConnectorErrorCode.CONFIGURATION)

    def check(name: str, count: int) -> None:
        if count > bounds[name]:
            _fail(
                "resource_limit",
                ConnectorErrorCode.RESOURCE_LIMIT_EXCEEDED,
                limit=name,
                bound=bounds[name],
                observed=count,
            )

    if not isinstance(data, bytes):
        _fail("invalid_snapshot", ConnectorErrorCode.CONFIGURATION)
    if (
        not isinstance(expected, Mapping)
        or set(expected) != {"schema", "profile", "sheets"}
        or expected["schema"] != "otc.xlsx-physical/1.0"
        or expected["profile"] != "literal-artifact/1.0"
    ):
        _fail("invalid_expectation", ConnectorErrorCode.CONFIGURATION)
    check("archive_bytes", len(data))
    parts: dict[str, bytes] = {}
    roots: dict[str, Any] = {}
    total = 0
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            check("zip_members", len(archive.infolist()))
            for info in archive.infolist():
                name = info.filename
                if name in parts or info.flag_bits & 1:
                    _fail("duplicate_or_encrypted_member")
                kind = _part_kind(name)
                check("member_bytes", info.file_size)
                chunks = []
                count = 0
                for chunk in _expanded_member(data, info):
                    count += len(chunk)
                    total += len(chunk)
                    check("member_bytes", count)
                    check("decompressed_bytes", total)
                    chunks.append(chunk)
                content = b"".join(chunks)
                parts[name] = content
                if kind != "image":
                    if (
                        b"<!DOCTYPE" in content.upper()
                        or b"<!ENTITY" in content.upper()
                        or b"\x00" in content
                    ):
                        _fail("unsafe_xml")
                    roots[name] = ET.fromstring(content)
    except ConnectorError:
        raise
    except (BadZipFile, ET.ParseError, RuntimeError, ValueError, OSError, EOFError):
        _fail("malformed_archive", ConnectorErrorCode.PROTOCOL_INVALID)
    required = {
        "[Content_Types].xml",
        "_rels/.rels",
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/styles.xml",
    }
    if not required <= parts.keys():
        _fail("missing_part")
    _content_types(parts, roots)
    _metadata_schema(roots)
    graph: dict[str, dict[str, tuple[str, str]]] = {}
    referenced = {"[Content_Types].xml", "_rels/.rels"}
    allowed = {
        "officeDocument": "workbook",
        "worksheet": "worksheet",
        "styles": "styles",
        "theme": "theme",
        "sharedStrings": "sharedStrings",
        "drawing": "drawing",
        "image": "image",
        "metadata/core-properties": "core",
        "extended-properties": "app",
    }
    for name, root in roots.items():
        if _part_kind(name) != "rels":
            continue
        owner = "" if name == "_rels/.rels" else name.replace("/_rels/", "/")[:-5]
        if owner and owner not in parts:
            _fail("orphan_relationships")
        relationships: dict[str, tuple[str, str]] = {}
        if root.tag != f"{{{P}}}Relationships":
            _fail("invalid_relationship_namespace")
        for rel in root:
            if (
                rel.tag != f"{{{P}}}Relationship"
                or set(rel.attrib) - {"Id", "Type", "Target", "TargetMode"}
                or rel.get("TargetMode", "Internal") != "Internal"
            ):
                _fail("external_or_invalid_relationship")
            rid, typ, target = rel.get("Id"), rel.get("Type", ""), rel.get("Target", "")
            if (
                not rid
                or rid in relationships
                or not target
                or "\\" in target
                or "?" in target
                or "#" in target
                or "%" in target
                or ":" in target
            ):
                _fail("invalid_relationship")
            if ".." in target.split("/"):
                _fail("escaping_relationship")
            resolved = (
                target.lstrip("/")
                if target.startswith("/")
                else posixpath.normpath(posixpath.join(posixpath.dirname(owner), target))
            )
            suffix = (
                typ.removeprefix(R + "/")
                if typ.startswith(R + "/")
                else (
                    "metadata/core-properties"
                    if typ
                    == "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties"
                    else ""
                )
            )
            if (
                suffix not in allowed
                or resolved not in parts
                or _part_kind(resolved) != allowed[suffix]
            ):
                _fail("invalid_relationship_target")
            owner_kind = _part_kind(owner) if owner else "package"
            if suffix not in {
                "package": {"officeDocument", "metadata/core-properties", "extended-properties"},
                "workbook": {"worksheet", "styles", "theme", "sharedStrings"},
                "worksheet": {"drawing"},
                "drawing": {"image"},
            }.get(owner_kind, set()):
                _fail("invalid_relationship_owner")
            if any(existing[1] == resolved for existing in relationships.values()):
                _fail("duplicate_relationship_target")
            relationships[rid] = (suffix, resolved)
            referenced.add(resolved)
        graph[owner] = relationships
        referenced.add(name)
    if set(parts) - referenced:
        _fail("orphan_part")
    workbook = roots["xl/workbook.xml"]
    if workbook.tag != f"{{{S}}}workbook":
        _fail("invalid_workbook_namespace")
    namespaces = {
        S,
        R,
        P,
        D,
        A,
        C,
        "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
        "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties",
        "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes",
        "http://purl.org/dc/elements/1.1/",
        "http://purl.org/dc/terms/",
        "http://purl.org/dc/dcmitype/",
        "http://www.w3.org/2001/XMLSchema-instance",
    }
    for root in roots.values():
        if any(
            not node.tag.startswith("{") or node.tag[1:].split("}", 1)[0] not in namespaces
            for node in root.iter()
        ):
            _fail("unsupported_xml_namespace")
        if any(
            node.tag in (f"{{{S}}}f", f"{{{S}}}calcPr", f"{{{S}}}calcChain") for node in root.iter()
        ):
            _fail("calculation_or_formula")
    shared: list[str] = []
    if "xl/sharedStrings.xml" in roots:
        for node in roots["xl/sharedStrings.xml"]:
            if node.tag != f"{{{S}}}si":
                _fail("invalid_shared_strings")
            shared.append("".join(t.text or "" for t in node.iter(f"{{{S}}}t")))
    sheets = list(workbook.findall(f"{{{S}}}sheets/{{{S}}}sheet"))
    check("sheets", len(sheets))
    expected_sheets = expected["sheets"]
    if (
        not isinstance(expected_sheets, (list, tuple))
        or not sheets
        or [s.get("name") for s in sheets] != [s["name"] for s in expected_sheets]
    ):
        _fail("sheet_coverage")
    if len({s.get("name", "").casefold() for s in sheets}) != len(sheets) or len(
        {s.get("sheetId") for s in sheets}
    ) != len(sheets):
        _fail("duplicate_sheet")
    raw_values = []
    image_count = cell_count = text_count = image_bytes = physical_cells = 0
    used_sheets = set()
    used_drawings = set()
    used_media = set()
    for sheet, intent in zip(sheets, expected_sheets, strict=True):
        if set(intent) - {"name", "cells", "config", "merges", "images", "styles"} or not {
            "name",
            "cells",
            "config",
            "merges",
            "images",
        } <= set(intent):
            _fail("invalid_expectation")
        rel = graph["xl/workbook.xml"].get(sheet.get(f"{{{R}}}id"))
        if rel is None or rel[0] != "worksheet" or rel[1] in used_sheets:
            _fail("sheet_relationship")
        path = rel[1]
        used_sheets.add(path)
        root = roots[path]
        if root.tag != f"{{{S}}}worksheet" or root.attrib:
            _fail("invalid_worksheet_namespace")
        allowed_sheet = {
            "sheetPr",
            "dimension",
            "sheetViews",
            "sheetFormatPr",
            "cols",
            "sheetData",
            "mergeCells",
            "printOptions",
            "pageMargins",
            "pageSetup",
            "headerFooter",
            "drawing",
        }
        if any(node.tag not in {f"{{{S}}}{tag}" for tag in allowed_sheet} for node in root):
            _fail("unsupported_worksheet_element")
        if len([node for node in root if node.tag == f"{{{S}}}sheetData"]) != 1:
            _fail("invalid_sheet_data")
        rows = root.findall(f"{{{S}}}sheetData/{{{S}}}row")
        for row in rows:
            if set(row.attrib) - {
                "r",
                "spans",
                "s",
                "customFormat",
                "ht",
                "hidden",
                "customHeight",
                "outlineLevel",
                "collapsed",
                "thickTop",
                "thickBot",
                "ph",
            } or any(node.tag != f"{{{S}}}c" for node in row):
                _fail("invalid_row_structure")
            row_number = row.get("r", "")
            _coord("A" + row_number)
            if any(_coord(node.get("r", ""))[0] != int(row_number) for node in row):
                _fail("cell_row_mismatch")
        row_ids = [row.get("r") for row in rows]
        if len(set(row_ids)) != len(row_ids):
            _fail("duplicate_row")
        raw_cells = root.findall(f"{{{S}}}sheetData/{{{S}}}row/{{{S}}}c")
        if len(raw_cells) != len(list(root.iter(f"{{{S}}}c"))):
            _fail("misplaced_cell")
        seen = set()
        values = {}
        for cell in root.findall(f"{{{S}}}sheetData/{{{S}}}row/{{{S}}}c"):
            if set(cell.attrib) - {"r", "s", "t"}:
                _fail("unsupported_cell_attribute")
            coord = cell.get("r", "")
            _coord(coord)
            if coord in seen:
                _fail("duplicate_cell")
            seen.add(coord)
            physical_cells += 1
            check("cells", physical_cells)
            if any(node.tag not in {f"{{{S}}}is", f"{{{S}}}v"} for node in cell) or len(cell) > 1:
                _fail("unsupported_cell_structure")
            kind = cell.get("t", "n")
            val = cell.find(f"{{{S}}}v")
            inline = cell.find(f"{{{S}}}is")
            if kind == "inlineStr":
                if inline is not None and (
                    any(node.tag != f"{{{S}}}t" for node in inline) or len(inline) > 1
                ):
                    _fail("unsupported_inline_string")
                value = (
                    "".join(t.text or "" for t in inline.iter(f"{{{S}}}t"))
                    if inline is not None
                    else ""
                )
            elif kind == "s" and val is not None:
                try:
                    index = int(val.text or "")
                    if index < 0:
                        raise ValueError
                    value = shared[index]
                except (ValueError, IndexError):
                    _fail("invalid_shared_string_reference")
            elif val is not None or inline is not None:
                _fail("nonliteral_cell")
            else:
                continue
            values[coord] = value
            cell_count += 1
            text_count += len(value.encode("utf-8"))
            check("cells", cell_count)
            check("text_bytes", text_count)
        wanted = intent["cells"]
        if values != {k: v["value"] for k, v in wanted.items()} or any(
            v.get("type") != "s"
            or v.get("format") != "@"
            or set(v) != {"value", "type", "format", "style"}
            for v in wanted.values()
        ):
            _fail("cell_coverage")
        merges = [m.get("ref", "") for m in root.findall(f"{{{S}}}mergeCells/{{{S}}}mergeCell")]
        if merges != list(intent["merges"]):
            _fail("merge_coverage")
        rectangles: list[tuple[tuple[int, int], tuple[int, int]]] = []
        for merge in merges:
            ends = merge.split(":")
            if len(ends) != 2:
                _fail("invalid_merge")
            start, end = map(_coord, ends)
            if start[0] > end[0] or start[1] > end[1]:
                _fail("invalid_merge")
            for a, b in rectangles:
                if start[0] <= b[0] and a[0] <= end[0] and start[1] <= b[1] and a[1] <= end[1]:
                    _fail("overlapping_merges")
            rectangles.append((start, end))
            for coord in values:
                r, c = _coord(coord)
                if start[0] <= r <= end[0] and start[1] <= c <= end[1] and (r, c) != start:
                    _fail("hidden_merged_value")
        images = []
        for drawing in root.findall(f"{{{S}}}drawing"):
            edge = graph.get(path, {}).get(drawing.get(f"{{{R}}}id"))
            if edge is None or edge[0] != "drawing" or edge[1] in used_drawings:
                _fail("drawing_relationship")
            used_drawings.add(edge[1])
            for anchor in roots[edge[1]]:
                if anchor.tag != f"{{{D}}}oneCellAnchor":
                    _fail("unsupported_image_anchor")
                if [node.tag for node in anchor] != [
                    f"{{{D}}}{tag}" for tag in ("from", "ext", "pic", "clientData")
                ]:
                    _fail("unsupported_drawing_structure")
                marker = anchor.find(f"{{{D}}}from")
                extent = anchor.find(f"{{{D}}}ext")
                blips = list(anchor.iter(f"{{{A}}}blip"))
                if marker is None or extent is None or len(blips) != 1:
                    _fail("invalid_image_anchor")
                try:
                    col, row = (
                        int(marker.findtext(f"{{{D}}}col")),
                        int(marker.findtext(f"{{{D}}}row")),
                    )
                    if int(marker.findtext(f"{{{D}}}colOff")) or int(
                        marker.findtext(f"{{{D}}}rowOff")
                    ):
                        _fail("image_offset")
                    width, height = int(extent.get("cx")), int(extent.get("cy"))
                except (ValueError, TypeError):
                    _fail("invalid_image_anchor")
                from openpyxl.utils.cell import get_column_letter

                coordinate = f"{get_column_letter(col + 1)}{row + 1}"
                _coord(coordinate)
                if width <= 0 or height <= 0:
                    _fail("invalid_image_extent")
                image_rel = graph.get(edge[1], {}).get(blips[0].get(f"{{{R}}}embed"))
                if image_rel is None or image_rel[0] != "image" or image_rel[1] in used_media:
                    _fail("image_relationship")
                media_path = image_rel[1]
                used_media.add(media_path)
                media = parts[media_path]
                image_count += 1
                image_bytes += len(media)
                check("images", image_count)
                check("image_bytes", len(media))
                check("total_image_bytes", image_bytes)
                from PIL import Image

                try:
                    with Image.open(io.BytesIO(media)) as decoded:
                        check("image_pixels", decoded.width * decoded.height)
                        mime = Image.MIME.get(decoded.format)
                        if mime not in ("image/png", "image/jpeg"):
                            _fail("unsupported_image")
                        decoded.verify()
                except ConnectorError:
                    raise
                except Exception:
                    _fail("invalid_image")
                images.append(
                    dict(
                        sha256="sha256:" + hashlib.sha256(media).hexdigest(),
                        mime_type=mime,
                        anchor=coordinate,
                        width_emu=width,
                        height_emu=height,
                    )
                )
        if images != list(intent["images"]):
            _fail("image_coverage")
        raw_values.append(values)
    if (
        used_sheets != {p for p in parts if _part_kind(p) == "worksheet"}
        or used_drawings != {p for p in parts if _part_kind(p) == "drawing"}
        or used_media != {p for p in parts if _part_kind(p) == "image"}
    ):
        _fail("unreferenced_content")
    try:
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
        try:
            if book.sheetnames != [s["name"] for s in expected_sheets]:
                _fail("decoded_sheet_coverage")
            for ws, intent, values in zip(
                book.worksheets, expected_sheets, raw_values, strict=True
            ):
                for coordinate, value in values.items():
                    cell = ws[coordinate]
                    if (
                        (cell.value != value and not (value == "" and cell.value is None))
                        or cell.data_type not in ("s", "inlineStr")
                        or cell.number_format != "@"
                    ):
                        _fail("decoded_cell_mismatch")
                    actual_style = _cell_style(cell)
                    if any(
                        k not in actual_style or actual_style[k] != v
                        for k, v in intent["cells"][coordinate]["style"].items()
                    ):
                        _fail("style_mismatch")
                for coordinate, declared in intent.get("styles", {}).items():
                    _coord(coordinate)
                    if set(declared) != {"style", "format"} or coordinate in intent["cells"]:
                        _fail("invalid_style_expectation")
                    cell = ws[coordinate]
                    if cell.value is not None or cell.number_format != declared["format"]:
                        _fail("style_only_cell_mismatch")
                    actual_style = _cell_style(cell)
                    if any(
                        key not in actual_style or actual_style[key] != value
                        for key, value in declared["style"].items()
                    ):
                        _fail("style_mismatch")
                actual_config = _sheet_config(ws)
                if any(
                    k not in actual_config or actual_config[k] != v
                    for k, v in intent["config"].items()
                ):
                    _fail("layout_mismatch")
                if sorted(str(m) for m in ws.merged_cells.ranges) != sorted(
                    intent["merges"]
                ) or len(ws._images) != len(intent["images"]):
                    _fail("decoded_object_coverage")
        finally:
            book.close()
    except ConnectorError:
        raise
    except Exception:
        _fail("decoded_read_failed", ConnectorErrorCode.PROTOCOL_INVALID)
    canonical = json.dumps(
        expected, sort_keys=True, separators=(",", ":"), ensure_ascii=False, allow_nan=False
    ).encode("utf-8")
    return dict(
        profile="literal-artifact/1.0",
        verification_level="xlsx-physical",
        semantic_hash="sha256:" + hashlib.sha256(canonical).hexdigest(),
        content_hash="sha256:" + hashlib.sha256(data).hexdigest(),
        cells=cell_count,
        images=image_count,
        checks=[
            "bounded_zip",
            "relationships",
            "raw_cells",
            "merges",
            "raw_images",
            "decoded_cells",
            "styles",
            "layout",
        ],
    )


def _content_types(parts: Mapping[str, bytes], roots: Mapping[str, Any]) -> None:
    root = roots["[Content_Types].xml"]
    if root.tag != f"{{{C}}}Types":
        _fail("invalid_content_types")
    defaults = {}
    overrides = {}
    for node in root:
        if node.tag == f"{{{C}}}Default":
            key = node.get("Extension")
            if key in defaults:
                _fail("duplicate_content_type")
            defaults[key] = node.get("ContentType")
        elif node.tag == f"{{{C}}}Override":
            key = node.get("PartName", "").lstrip("/")
            if key in overrides or key not in parts:
                _fail("invalid_content_type")
            overrides[key] = node.get("ContentType")
        else:
            _fail("invalid_content_type")
    prefix = "application/vnd.openxmlformats-officedocument."
    allowed = dict(
        workbook=prefix + "spreadsheetml.sheet.main+xml",
        worksheet=prefix + "spreadsheetml.worksheet+xml",
        styles=prefix + "spreadsheetml.styles+xml",
        sharedStrings=prefix + "spreadsheetml.sharedStrings+xml",
        drawing=prefix + "drawing+xml",
        theme=prefix + "theme+xml",
        app=prefix + "extended-properties+xml",
        core="application/vnd.openxmlformats-package.core-properties+xml",
        rels="application/vnd.openxmlformats-package.relationships+xml",
    )
    for name in parts:
        kind = _part_kind(name)
        if kind == "types":
            continue
        actual = overrides.get(name, defaults.get(name.rsplit(".", 1)[-1]))
        wanted = (
            ("image/png" if name.endswith(".png") else "image/jpeg")
            if kind == "image"
            else allowed[kind]
        )
        if actual != wanted:
            _fail("unsupported_content_type")


def read_archive(data: bytes, limits: Any = None) -> dict[str, bytes]:
    """Bound general XLSX expansion/XML before a provider decoder is invoked."""
    bounds = dict(_DEFAULTS)
    if limits is not None:
        supplied = (
            dict(limits)
            if isinstance(limits, Mapping)
            else {k: getattr(limits, k) for k in bounds if hasattr(limits, k)}
        )
        if set(supplied) - set(bounds):
            _fail("invalid_limits", ConnectorErrorCode.CONFIGURATION)
        bounds.update(supplied)
    if any(type(v) is not int or v <= 0 for v in bounds.values()):
        _fail("invalid_limits", ConnectorErrorCode.CONFIGURATION)

    def check(name: str, count: int) -> None:
        if count > bounds[name]:
            _fail(
                "resource_limit",
                ConnectorErrorCode.RESOURCE_LIMIT_EXCEEDED,
                limit=name,
                bound=bounds[name],
                observed=count,
            )

    check("archive_bytes", len(data))
    result = {}
    total = 0
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            check("zip_members", len(archive.infolist()))
            for info in archive.infolist():
                name = info.filename
                if (
                    name in result
                    or name.startswith("/")
                    or ".." in name.split("/")
                    or "\\" in name
                    or info.flag_bits & 1
                ):
                    _fail("invalid_archive_member")
                check("member_bytes", info.file_size)
                content = bytearray()
                for chunk in _expanded_member(data, info):
                    total += len(chunk)
                    check("member_bytes", len(content) + len(chunk))
                    check("decompressed_bytes", total)
                    content.extend(chunk)
                value = bytes(content)
                if name.endswith((".xml", ".rels")):
                    if (
                        b"<!DOCTYPE" in value.upper()
                        or b"<!ENTITY" in value.upper()
                        or b"\x00" in value
                    ):
                        _fail("unsafe_xml")
                    ET.fromstring(value)
                result[name] = value
    except ConnectorError:
        raise
    except (BadZipFile, ET.ParseError, RuntimeError, ValueError, OSError, EOFError):
        _fail("malformed_archive", ConnectorErrorCode.PROTOCOL_INVALID)
    return result


def verify_snapshot(
    data: bytes, expected: Mapping[str, Any], limits: Any = None
) -> Mapping[str, Any]:
    """Verify one immutable byte snapshot without exposing parser/input contents."""
    try:

        def plain(value: Any) -> Any:
            if isinstance(value, Mapping):
                if any(not isinstance(key, str) for key in value):
                    _fail("invalid_expectation")
                return {key: plain(item) for key, item in value.items()}
            if isinstance(value, (list, tuple)):
                return [plain(item) for item in value]
            if value is None or type(value) in (str, int, bool):
                return value
            if type(value) is float and math.isfinite(value):
                return value
            _fail("invalid_expectation")

        normalized = plain(expected)
        _expectation_schema(normalized)
        return _verify_snapshot(data, normalized, limits)
    except ConnectorError:
        raise
    except (KeyError, TypeError, ValueError, AttributeError, OverflowError):
        _fail("invalid_expectation_or_structure", ConnectorErrorCode.PROTOCOL_INVALID)


def _metadata_schema(roots: Mapping[str, Any]) -> None:
    """Close metadata/workbook shapes that decoders otherwise silently ignore."""
    core_ns = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
    app_ns = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    core_fields = (
        {
            f"{{{core_ns}}}{name}"
            for name in (
                "category",
                "contentStatus",
                "keywords",
                "lastModifiedBy",
                "lastPrinted",
                "revision",
                "version",
            )
        }
        | {
            f"{{http://purl.org/dc/elements/1.1/}}{name}"
            for name in ("creator", "description", "identifier", "language", "subject", "title")
        }
        | {f"{{http://purl.org/dc/terms/}}{name}" for name in ("created", "modified")}
    )
    if "docProps/core.xml" in roots:
        root = roots["docProps/core.xml"]
        if (
            root.tag != f"{{{core_ns}}}coreProperties"
            or root.attrib
            or any(node.tag not in core_fields or len(node) for node in root)
        ):
            _fail("unsupported_core_metadata")
    if "docProps/app.xml" in roots:
        root = roots["docProps/app.xml"]
        fields = {
            f"{{{app_ns}}}{name}"
            for name in (
                "Application",
                "AppVersion",
                "DocSecurity",
                "ScaleCrop",
                "Company",
                "LinksUpToDate",
                "SharedDoc",
                "HyperlinksChanged",
            )
        }
        if (
            root.tag != f"{{{app_ns}}}Properties"
            or root.attrib
            or any(node.tag not in fields or len(node) or node.attrib for node in root)
        ):
            _fail("unsupported_app_metadata")
    root = roots["xl/workbook.xml"]
    allowed = {
        f"{{{S}}}{name}"
        for name in ("workbookPr", "workbookProtection", "bookViews", "sheets", "definedNames")
    }
    if (
        root.attrib
        or any(node.tag not in allowed for node in root)
        or len({node.tag for node in root}) != len(root)
    ):
        _fail("unsupported_workbook_structure")
    for node in root.findall(f"{{{S}}}sheets/{{{S}}}sheet"):
        if set(node.attrib) - {"name", "sheetId", "state", f"{{{R}}}id"} or len(node):
            _fail("unsupported_sheet_metadata")
    for node in root.findall(f"{{{S}}}definedNames/{{{S}}}definedName"):
        if (
            node.get("name") != "_xlnm.Print_Area"
            or set(node.attrib) - {"name", "localSheetId"}
            or len(node)
        ):
            _fail("unsupported_defined_name")


def _expectation_schema(expected: Any) -> None:
    if not isinstance(expected, dict) or not isinstance(expected.get("sheets"), list):
        _fail("invalid_expectation")

    def style_check(style: Any) -> None:
        if not isinstance(style, dict):
            _fail("invalid_style_expectation")
        for key, value in style.items():
            if key in ("bold", "italic", "wrap_text"):
                valid = type(value) is bool
            elif key == "size":
                valid = value is None or type(value) in (int, float) and 0 < value <= 409
            elif key in ("family", "foreground", "fill", "vertical", "horizontal"):
                valid = value is None or isinstance(value, str)
            elif key == "border":
                valid = isinstance(value, dict) and not set(value) - {
                    "left",
                    "right",
                    "top",
                    "bottom",
                }
                if valid:
                    for edge in value.values():
                        if (
                            not isinstance(edge, dict)
                            or set(edge) != {"style", "color"}
                            or edge["style"]
                            not in ("thin", "medium", "thick", "dashed", "dotted", "double")
                        ):
                            _fail("invalid_border_expectation")
            else:
                valid = False
            if not valid:
                _fail("invalid_style_expectation")

    for sheet in expected["sheets"]:
        if not isinstance(sheet, dict):
            _fail("invalid_expectation")
        for cell in sheet.get("cells", {}).values():
            style_check(cell.get("style"))
        for cell in sheet.get("styles", {}).values():
            style_check(cell.get("style"))
        for picture in sheet.get("images", []):
            if not isinstance(picture, dict) or set(picture) != {
                "sha256",
                "mime_type",
                "anchor",
                "width_emu",
                "height_emu",
            }:
                _fail("invalid_image_expectation")
            if (
                any(
                    type(picture[key]) is not int or picture[key] <= 0
                    for key in ("width_emu", "height_emu")
                )
                or picture["mime_type"] not in ("image/png", "image/jpeg")
                or not isinstance(picture["sha256"], str)
                or not re.fullmatch(r"sha256:[0-9a-f]{64}", picture["sha256"])
            ):
                _fail("invalid_image_expectation")
            _coord(picture["anchor"])


def _expanded_member(data: bytes, info: ZipInfo) -> Iterator[bytes]:
    """Expand independently of the advertised output size, with bounded chunks."""
    offset = info.header_offset
    if offset < 0 or data[offset : offset + 4] != b"PK\x03\x04" or offset + 30 > len(data):
        _fail("malformed_archive")
    flags, method = struct.unpack_from("<HH", data, offset + 6)
    name_size, extra_size = struct.unpack_from("<HH", data, offset + 26)
    start = offset + 30 + name_size + extra_size
    end = start + info.compress_size
    if flags & 1 or method != info.compress_type or method not in (0, 8) or end > len(data):
        _fail("unsupported_compression")
    try:
        name = data[offset + 30 : offset + 30 + name_size].decode(
            "utf-8" if flags & 0x800 else "cp437"
        )
    except UnicodeDecodeError:
        _fail("malformed_archive")
    if name != info.orig_filename:
        _fail("member_name_mismatch")
    decoder = zlib.decompressobj(-15) if method == 8 else None
    count = crc = 0
    for position in range(start, end, 65536):
        pending = data[position : min(position + 65536, end)]
        while pending:
            try:
                expanded = decoder.decompress(pending, 65536) if decoder is not None else pending
            except zlib.error:
                _fail("malformed_deflate")
            pending = decoder.unconsumed_tail if decoder is not None else b""
            if decoder is not None and decoder.unused_data:
                _fail("trailing_compressed_data")
            count += len(expanded)
            crc = zlib.crc32(expanded, crc)
            if expanded:
                yield expanded
    if decoder is not None and not decoder.eof:
        _fail("incomplete_deflate")
    if count != info.file_size or crc != info.CRC:
        _fail("member_size_or_crc_mismatch")
