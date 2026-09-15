from pathlib import Path

import pytest
from open_table_connector.contract import CapabilityIdentity, ConnectorError
from open_table_connector.spreadsheets import SpreadsheetTarget
from open_table_connector.spreadsheets._operations import Change


def change(operation, sheet="", **arguments):
    return Change(
        operation, CapabilityIdentity(f"spreadsheet.{operation}", "1.0"), sheet, arguments
    )


def binding(provider, path, *, new=True, profile="general/1.0"):
    result = provider.bind(SpreadsheetTarget(path.as_uri()))
    result.update(new=new, profile=profile)
    return result


def test_local_buffer_preview_and_commit(tmp_path):
    from open_table_connector.local_files.spreadsheet_workbook import LocalSpreadsheetProvider

    provider = LocalSpreadsheetProvider()
    path = tmp_path / "book.xlsx"
    bound = binding(provider, path)
    changes = (
        change("worksheet.create", "Report", name="Report"),
        change("range.write", "Report", address="A1:B1", values=[[1, "=literal"]]),
    )
    provider.preflight(bound, changes)
    assert not path.exists()
    preview = provider.observe(
        bound,
        {"operation": "range.read", "target_key": "Report", "address": "A1:B1", "changes": changes},
    )
    assert preview["value"] == [[1, "=literal"]]
    result = provider.commit(
        bound, changes, allow_partial=False, expected_revision=None, idempotency_key=None
    )
    assert result["commit"] == "committed"
    from openpyxl import load_workbook

    book = load_workbook(path)
    assert book["Report"]["B1"].data_type == "s"
    assert book["Report"]["A1"].value == 1
    book.close()


def test_competing_creator_cannot_delete_winner(tmp_path):
    from open_table_connector.local_files.spreadsheet_workbook import LocalSpreadsheetProvider

    provider = LocalSpreadsheetProvider()
    path = tmp_path / "book.xlsx"
    bound = binding(provider, path)
    path.write_bytes(b"winner")
    with pytest.raises(ConnectorError):
        provider.commit(
            bound,
            (change("worksheet.create", "Report", name="Report"),),
            allow_partial=False,
            expected_revision=None,
            idempotency_key=None,
        )
    assert path.read_bytes() == b"winner"


def test_ragged_matrix_and_literal_formula_rejected_before_mutation(tmp_path):
    from open_table_connector.local_files.spreadsheet_workbook import LocalSpreadsheetProvider

    provider = LocalSpreadsheetProvider()
    bound = binding(provider, tmp_path / "book.xlsx", profile="literal-artifact/1.0")
    first = change("worksheet.create", "Report", name="Report")
    with pytest.raises(ConnectorError):
        provider.preflight(
            bound,
            (first, change("range.write", "Report", address="A1:B2", values=[["a", "b"], ["c"]])),
        )
    with pytest.raises(ConnectorError):
        provider.preflight(
            bound, (first, change("formula.set", "Report", address="A1", expression="=1+1"))
        )


def client():
    from open_table_connector.local_files import LocalFilesConnector
    from open_table_connector.sdk import Client, ConnectorRegistry

    return Client(registry=ConnectorRegistry([LocalFilesConnector()]))


def test_literal_layout_images_empty_and_receipt_roundtrip(tmp_path):
    from io import BytesIO

    from open_table_connector.spreadsheets import ImageSpec
    from PIL import Image

    path = tmp_path / "literal.xlsx"
    book = client().workbook.create(path.as_uri())
    sheet = book.worksheet.create("Report")
    sheet.range("A1:B2").write([["title", ""], ["=literal", "42"]])
    sheet.range("A1:B2").style(
        family="Arial",
        size=10,
        bold=True,
        foreground="#183245",
        fill="#F0F5F8",
        vertical="top",
        wrap_text=True,
        border={"bottom": {"style": "thin", "color": "#183245"}},
    )
    sheet.config(
        row_heights={"1": 28},
        column_widths={"A": 32, "B": 29},
        show_gridlines=False,
        freeze_panes="A2",
        print_area="A1:B8",
        orientation="landscape",
        paper_size="9",
        fit_to_page=True,
        fit_width=1,
        fit_height=0,
        horizontal_centered=True,
        margins={
            "left": 0.3,
            "right": 0.3,
            "top": 0.3,
            "bottom": 0.3,
            "header": 0.1,
            "footer": 0.1,
        },
    )
    stream = BytesIO()
    Image.new("RGB", (4, 3), "blue").save(stream, format="PNG")
    sheet.image(
        ImageSpec(content=stream.getvalue(), mime_type="image/png", anchor="A5"),
        width=40,
        height=30,
    )
    result = book.write()
    assert result.commit.value == "committed"
    assert book.verify().verification.value == "passed"
    # Existing receipt serialization retains independent intent.
    wire = result.to_wire()
    assert (
        wire["receipts"][0]["details"]["expected"]["sheets"][0]["cells"]["A2"]["value"]
        == "=literal"
    )
    from open_table_connector.sdk import OTCError

    reopened = client().workbook(path.as_uri(), profile="literal-artifact/1.0")
    with pytest.raises(OTCError):
        reopened.verify()
    assert (
        reopened.verify(expected=wire["receipts"][0]["details"]["expected"]).verification.value
        == "passed"
    )


def test_failure_before_publish_keeps_quarantine(tmp_path, monkeypatch):
    import open_table_connector.local_files.spreadsheet_verify as verifier
    from open_table_connector.sdk import OTCError

    path = tmp_path / "failed.xlsx"
    quarantine = tmp_path / "quarantine"
    book = client().workbook.create(path.as_uri(), failure_directory=quarantine)
    book.worksheet.create("R").range("A1").write([["safe"]])

    def fail(*args, **kwargs):
        raise ConnectorError.configuration("injected verification failure")

    monkeypatch.setattr(verifier, "verify_snapshot", fail)
    with pytest.raises(OTCError) as raised:
        book.write()
    assert not path.exists()
    evidence = raised.value.result.error.safe_details["evidence"]
    assert Path(evidence["path"]).is_file()
    assert evidence["complete"] is True
    assert raised.value.result.commit.value == "not_committed"


def test_postcommit_cleanup_failure_keeps_destination_and_freezes(tmp_path, monkeypatch):
    from open_table_connector.sdk import OTCError

    path = tmp_path / "committed.xlsx"
    book = client().workbook.create(path.as_uri())
    book.worksheet.create("R").range("A1").write([["safe"]])
    original = Path.unlink

    def fail(stage, *args, **kwargs):
        if stage.name.startswith(".committed-"):
            raise OSError("injected cleanup")
        return original(stage, *args, **kwargs)

    monkeypatch.setattr(Path, "unlink", fail)
    with pytest.raises(OTCError) as raised:
        book.write()
    assert path.exists()
    assert raised.value.result.commit.value == "committed"
    with pytest.raises(OTCError):
        book.worksheet("R").range("A1").write([["retry"]])


def test_general_edit_preserves_objects_and_literal_strings(tmp_path):
    from openpyxl import Workbook, load_workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo

    path = tmp_path / "objects.xlsx"
    original = Workbook()
    sheet = original.active
    sheet.title = "Inputs"
    sheet.append(["Name", "Value"])
    sheet.append(["a", 1])
    sheet.append(["b", 2])
    sheet["D1"] = "=SUM(B2:B3)"
    table = Table(displayName="Data", ref="A1:B3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2")
    sheet.add_table(table)
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    sheet.add_chart(chart, "F1")
    original.defined_names.add(DefinedName("Amount", attr_text="'Inputs'!$B$2"))
    validation = DataValidation(type="whole", operator="between", formula1=0, formula2=100)
    sheet.add_data_validation(validation)
    validation.add("B2")
    hidden = original.create_sheet("Hidden")
    hidden.sheet_state = "veryHidden"
    hidden["A1"] = "secret"
    hidden.protection.sheet = True
    original.save(path)
    original.close()
    book = client().workbook(path.as_uri())
    book.worksheet("Inputs").range("B2").write([[42]])
    book.write()
    book.worksheet("Inputs").range("E1").write([["=literal"]])
    book.write()
    actual = load_workbook(path)
    assert actual["Inputs"]["B2"].value == 42
    assert actual["Inputs"]["D1"].value == "=SUM(B2:B3)"
    assert actual["Inputs"]["E1"].data_type == "s"
    assert len(actual["Inputs"]._charts) == 1 and "Data" in actual["Inputs"].tables
    assert (
        "Amount" in actual.defined_names
        and len(actual["Inputs"].data_validations.dataValidation) == 1
    )
    assert actual["Hidden"].sheet_state == "veryHidden" and actual["Hidden"].protection.sheet
    actual.close()


def test_stale_session_and_symlink_fail_without_modification(tmp_path):
    from open_table_connector.sdk import OTCError

    path = tmp_path / "stale.xlsx"
    book = client().workbook.create(path.as_uri())
    book.worksheet.create("R").range("A1").write([["old"]])
    book.write()
    first = client().workbook(path.as_uri())
    second = client().workbook(path.as_uri())
    first.worksheet("R").range("A1").write([["first"]])
    second.worksheet("R").range("A1").write([["second"]])
    first.write()
    before = path.read_bytes()
    with pytest.raises(OTCError):
        second.write()
    assert path.read_bytes() == before
    alias = tmp_path / "alias.xlsx"
    alias.symlink_to(path)
    with pytest.raises(OTCError):
        client().workbook(alias.as_uri())
    assert path.read_bytes() == before


def test_two_creators_have_exactly_one_winner(tmp_path):
    from concurrent.futures import ThreadPoolExecutor

    from open_table_connector.sdk import OTCError

    path = tmp_path / "race.xlsx"
    books = [client().workbook.create(path.as_uri()) for _ in range(2)]
    for index, book in enumerate(books):
        book.worksheet.create("R").range("A1").write([[str(index)]])

    def publish(book):
        try:
            return book.write().commit.value
        except OTCError:
            return "rejected"

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(publish, books))
    assert sorted(results) == ["committed", "rejected"]
    assert path.is_file()


def test_general_dates_empty_styles_and_formulas(tmp_path):
    from datetime import date

    path = tmp_path / "dates.xlsx"
    book = client().workbook.create(path.as_uri(), profile="general/1.0")
    sheet = book.worksheet.create("R")
    sheet.range("A1:B2").write([[date(2025, 1, 1), ""], [60, None]])
    sheet.range("B1").style(bold=True)
    sheet.formulas().set("C1", "=A2+1")
    sheet.range("D1").style(italic=True)
    book.write()
    assert book.verify().verification.value == "passed"


def test_style_only_literal_cells_are_verified(tmp_path):
    from open_table_connector.sdk import OTCError
    from openpyxl import load_workbook

    path = tmp_path / "styles.xlsx"
    book = client().workbook.create(path.as_uri())
    sheet = book.worksheet.create("R")
    sheet.range("A1").write([["safe"]])
    sheet.range("D9").style(bold=True)
    book.write()
    changed = load_workbook(path)
    from copy import copy

    font = copy(changed["R"]["D9"].font)
    font.bold = False
    changed["R"]["D9"].font = font
    changed.calculation = None
    changed.save(path)
    changed.close()
    with pytest.raises(OTCError):
        book.verify()


def test_failure_retention_failure_is_reported(tmp_path, monkeypatch):
    import open_table_connector.local_files.spreadsheet_verify as verifier
    from open_table_connector.sdk import OTCError

    path = tmp_path / "failed.xlsx"
    bad_dir = tmp_path / "file"
    bad_dir.write_text("preserve")
    book = client().workbook.create(path.as_uri(), failure_directory=bad_dir)
    book.worksheet.create("R").range("A1").write([["x"]])

    def fail(*args, **kwargs):
        raise ConnectorError.configuration("primary")

    monkeypatch.setattr(verifier, "verify_snapshot", fail)
    with pytest.raises(OTCError) as error:
        book.write()
    assert error.value.result.error.safe_details["retention_error"] == "FileExistsError"
    assert bad_dir.read_text() == "preserve" and not path.exists()


def test_formula_extension_changes_invalidate_workbook_session(tmp_path):
    from open_table_connector.sdk import OTCError

    path = tmp_path / "formula-race.xlsx"
    book = client().workbook.create(path.as_uri(), profile="general/1.0")
    book.worksheet.create("R").range("A1").write([[1]])
    book.write()
    opened = client().workbook(path.as_uri())
    opened.worksheet("R").range("A1").write([[2]])
    # A cooperating immediate writer uses the same .otc.lock file; observation hash protects stale sessions.
    from openpyxl import load_workbook

    other = load_workbook(path)
    other["R"]["B1"] = "=1+1"
    other.save(path)
    other.close()
    with pytest.raises(OTCError):
        opened.write()


def test_reject_rich_text_before_general_edit(tmp_path):
    from io import BytesIO
    from zipfile import ZIP_DEFLATED, ZipFile

    from open_table_connector.sdk import OTCError
    from openpyxl import Workbook

    path = tmp_path / "rich.xlsx"
    book = Workbook()
    book.active["A1"] = "rich"
    book.save(path)
    book.close()
    buf = BytesIO()
    with ZipFile(path) as original, ZipFile(buf, "w", ZIP_DEFLATED) as changed:
        for name in original.namelist():
            data = original.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = data.replace(
                    b"<is><t>rich</t></is>", b"<is><r><rPr><b/></rPr><t>rich</t></r></is>"
                )
            changed.writestr(name, data)
    path.write_bytes(buf.getvalue())
    before = path.read_bytes()
    with pytest.raises(OTCError):
        client().workbook(path.as_uri())
    assert path.read_bytes() == before


def test_sheet_rename_rejects_validation_references(tmp_path):
    from open_table_connector.sdk import OTCError
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation

    path = tmp_path / "refs.xlsx"
    book = Workbook()
    book.active.title = "Inputs"
    book.active["A1"] = "x"
    dv = DataValidation(type="list", formula1="'Inputs'!$A$1:$A$2")
    book.active.add_data_validation(dv)
    dv.add("B1")
    book.save(path)
    book.close()
    session = client().workbook(path.as_uri())
    with pytest.raises(OTCError):
        session.worksheet("Inputs").rename("Changed")


def test_default_reopen_verifies_explicit_strict_manifest(tmp_path):
    path = tmp_path / "verified.xlsx"
    book = client().workbook.create(path.as_uri())
    sheet = book.worksheet.create("R")
    sheet.range("A1").write([["x"]])
    sheet.range("A1:B1").merge()
    sheet.range("A3:B3").merge()
    result = book.write()
    expected = result.to_wire()["receipts"][0]["details"]["expected"]
    assert client().workbook(path.as_uri()).verify(expected=expected).verification.value == "passed"


def test_sort_rejects_hyperlinks_before_values_detach(tmp_path):
    from open_table_connector.sdk import OTCError
    from openpyxl import Workbook, load_workbook

    path = tmp_path / "linked-sort.xlsx"
    original = Workbook()
    sheet = original.active
    sheet["A1"] = "b"
    sheet["A1"].hyperlink = "https://example.com/b"
    sheet["A2"] = "a"
    sheet["A2"].hyperlink = "https://example.com/a"
    original.save(path)
    original.close()
    before = path.read_bytes()

    session = client().workbook(path.as_uri())
    with pytest.raises(OTCError) as error:
        session.worksheet("Sheet").range("A1:A2").sort()
    assert error.value.result.error.code.value == "unsupported_capability"
    assert path.read_bytes() == before
    assert session.worksheet("Sheet").range("A1:A2").read().value == [["b"], ["a"]]
    observed = load_workbook(path)
    try:
        assert observed.active["A1"].hyperlink.target == "https://example.com/b"
        assert observed.active["A2"].hyperlink.target == "https://example.com/a"
    finally:
        observed.close()
        session.close()


@pytest.mark.parametrize("profile", ["general/1.0", "literal-artifact/1.0"])
def test_reopened_literal_empty_string_remains_distinct_from_blank(tmp_path, profile):
    path = tmp_path / "empty-literal.xlsx"
    book = client().workbook.create(path.as_uri())
    book.worksheet.create("Report").range("A1").write([[""]])
    receipt = book.write()
    book.close()

    reopened = client().workbook(path.as_uri(), profile=profile)
    try:
        assert reopened.worksheet("Report").range("A1:B1").read().value == [["", None]]
        expected = receipt.to_wire()["receipts"][0]["details"]["expected"]
        assert reopened.verify(expected=expected).verification.value == "passed"
    finally:
        reopened.close()


@pytest.mark.parametrize(
    "part", ["xl/pivotTables/pivotTable1.xml", "xl/pivotCache/pivotCacheRecords1.xml"]
)
def test_pivot_parts_rejected_before_unproven_preservation(part):
    from io import BytesIO
    from zipfile import ZipFile

    from open_table_connector.local_files.spreadsheet_workbook import _preservation
    from open_table_connector.spreadsheets._limits import ArtifactLimits

    archive = BytesIO()
    with ZipFile(archive, "w") as container:
        container.writestr(
            part,
            '<pivotCacheRecords xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>',
        )
    with pytest.raises(ConnectorError) as error:
        _preservation(archive.getvalue(), ArtifactLimits())
    assert error.value.safe_details["reason"] == "unsupported_part"
