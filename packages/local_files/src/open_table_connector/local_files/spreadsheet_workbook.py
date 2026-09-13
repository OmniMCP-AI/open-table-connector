"""Unified local workbook operations for Excel-compatible XLSX files.

The session deliberately keeps the public surface small.  It is also the
reference implementation used by the local-files provider for the
``literal-artifact/1.0`` profile; provider adapters can expose the same
operations without importing this module.
"""

from __future__ import annotations

import hashlib
import json
import os
import posixpath
import re
import tempfile
from collections.abc import Iterable, Mapping, Sequence
from copy import copy
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from open_table_connector.contract import CapabilityIdentity, TableURI
from open_table_connector.sdk.model import TableMode
from open_table_connector.sdk.result import (
    CommitState,
    ErrorCode,
    ErrorInfo,
    OperationResult,
    OTCError,
    Outcome,
    Receipt,
    VerificationState,
)
from open_table_connector.spreadsheets import (
    ArtifactLimits,
    CellFormat,
    CellStyle,
    ImageSpec,
    RangeRef,
)

_A1 = re.compile(r"^([A-Z]{1,3})([1-9][0-9]*)$")
_CELL_LIMIT = 32_767
_PROFILE = "literal-artifact/1.0"
_WORKBOOK_WRITE = CapabilityIdentity("spreadsheet.workbook.write", "1.0")
_WORKBOOK_VERIFY = CapabilityIdentity("spreadsheet.workbook.verify", "1.0")
_RANGE_READ = CapabilityIdentity("spreadsheet.range.read", "1.0")
_RANGE_WRITE = CapabilityIdentity("spreadsheet.range.write", "1.0")
_RANGE_STYLE = CapabilityIdentity("spreadsheet.range.style", "1.0")
_RANGE_FORMAT = CapabilityIdentity("spreadsheet.range.format", "1.0")
_RANGE_SORT = CapabilityIdentity("spreadsheet.range.sort", "1.0")
_IMAGE_INSERT = CapabilityIdentity("spreadsheet.image.insert", "1.0")


def _failure(message: str, code: ErrorCode, **details: object) -> OTCError:
    result = OperationResult[None](
        value=None,
        outcome=Outcome.REJECTED,
        commit=CommitState.NOT_STARTED,
        verification=VerificationState.SKIPPED,
        receipts=(),
        error=ErrorInfo(code=code, message=message, safe_details=details),
    )
    return OTCError(message, result)


def _path_from_uri(value: str | TableURI) -> tuple[TableURI, Path]:
    uri = value if isinstance(value, TableURI) else TableURI(value)
    parsed = urlsplit(uri.value)
    if parsed.scheme != "file" or parsed.netloc not in {"", "localhost"} or parsed.query or parsed.fragment:
        raise _failure("workbook target must be an absolute file URI without query or fragment", ErrorCode.INVALID_TARGET)
    path = Path(unquote(parsed.path))
    if not path.is_absolute() or path.suffix.casefold() != ".xlsx":
        raise _failure("workbook target must be an absolute .xlsx file URI", ErrorCode.INVALID_TARGET)
    return uri, path


def _column_number(value: str) -> int:
    number = 0
    for char in value:
        number = number * 26 + ord(char) - 64
    return number


def _coordinates(address: str) -> tuple[int, int, int, int]:
    match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]*)(?::([A-Z]{1,3})([1-9][0-9]*))?", address.upper())
    if match is None:
        raise _failure("range must be a finite A1 rectangle", ErrorCode.INVALID_TARGET)
    first_col, first_row, last_col, last_row = match.groups()
    c1, r1 = _column_number(first_col), int(first_row)
    c2, r2 = _column_number(last_col or first_col), int(last_row or first_row)
    if c1 > 16_384 or c2 > 16_384 or r1 > 1_048_576 or r2 > 1_048_576 or c2 < c1 or r2 < r1:
        raise _failure("range is outside the Excel worksheet bounds", ErrorCode.INVALID_TARGET)
    return r1, c1, r2, c2


def _literal(cell: Any) -> str:
    if isinstance(cell, str):
        text = cell
    elif cell is None:
        text = ""
    else:
        text = str(cell)
    if len(text) > _CELL_LIMIT:
        raise _failure("literal cell text exceeds Excel's 32,767 character limit", ErrorCode.RESOURCE_LIMIT)
    return text


def _safe_json(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Mapping):
        return {str(key): _safe_json(item) for key, item in value.items()}
    if isinstance(value, Sequence) and not isinstance(value, (str, bytes, bytearray)):
        return [_safe_json(item) for item in value]
    return str(value)


def _result(value: Any, capability: CapabilityIdentity, uri: TableURI, *, operation: str, details: Mapping[str, Any] = ()) -> OperationResult[Any]:
    receipt = Receipt(
        kind="spreadsheet",
        operation=operation,
        connector_id="local-files",
        capability=capability.to_reference(),
        safe_target=uri,
        mode=TableMode.SHEET_MODE,
        details=dict(details),
    )
    return OperationResult(
        value=value,
        outcome=Outcome.SUCCEEDED,
        commit=CommitState.COMMITTED,
        verification=VerificationState.PASSED,
        receipts=(receipt,),
    )


@dataclass(frozen=True, slots=True)
class FormulaWrite:
    address: str
    expression: str
    result: Any = None


@dataclass(frozen=True, slots=True)
class WorkbookVerification:
    status: str
    content_hash: str
    semantic_hash: str
    sheets: tuple[str, ...]
    cells: int
    formulas: int
    images: int

    def to_wire(self) -> dict[str, Any]:
        return {
            "status": self.status,
            "content_hash": self.content_hash,
            "semantic_hash": self.semantic_hash,
            "sheets": list(self.sheets),
            "cells": self.cells,
            "formulas": self.formulas,
            "images": self.images,
        }


@dataclass(slots=True)
class WorkbookSession:
    """In-memory workbook handle with explicit write and verify operations."""

    uri: TableURI
    profile: str = _PROFILE
    limits: ArtifactLimits = field(default_factory=ArtifactLimits)
    _book: Any = field(default=None, repr=False)
    _expected: dict[str, dict[str, dict[str, Any]]] = field(default_factory=dict, repr=False)
    _images: list[dict[str, Any]] = field(default_factory=list, repr=False)
    _opened: bool = field(default=False, repr=False)
    _written: bool = field(default=False, repr=False)

    @classmethod
    def create(cls, uri: str | TableURI, *, profile: str = _PROFILE, limits: ArtifactLimits | None = None) -> WorkbookSession:
        target, path = _path_from_uri(uri)
        if profile != _PROFILE:
            raise _failure("unsupported workbook profile", ErrorCode.UNSUPPORTED_MODE, profile=profile)
        if path.exists():
            raise _failure("destination already exists; create is exclusive", ErrorCode.DESTINATION_EXISTS, path=str(path))
        try:
            from openpyxl import Workbook
        except ImportError:
            raise _failure("Excel workbook operations require openpyxl", ErrorCode.UNSUPPORTED_CAPABILITY) from None
        book = Workbook()
        book.remove(book.active)
        book.calculation.fullCalcOnLoad = False
        book.calculation.forceFullCalc = False
        book.calculation.calcMode = "manual"
        return cls(target, profile, limits or ArtifactLimits(), book)

    @classmethod
    def open(cls, uri: str | TableURI, *, limits: ArtifactLimits | None = None) -> WorkbookSession:
        target, path = _path_from_uri(uri)
        if not path.is_file():
            raise _failure("workbook file does not exist", ErrorCode.TARGET_NOT_FOUND, path=str(path))
        if path.stat().st_size > (limits or ArtifactLimits()).archive_bytes:
            raise _failure("workbook exceeds archive byte limit", ErrorCode.RESOURCE_LIMIT)
        try:
            from openpyxl import load_workbook
            book = load_workbook(path, data_only=False)
        except Exception as exc:
            raise _failure("workbook could not be opened", ErrorCode.ARTIFACT_INTEGRITY, reason=type(exc).__name__) from None
        session = cls(target, "general/1.0", limits or ArtifactLimits(), book, _opened=True, _written=True)
        session._capture_expected()
        return session

    @property
    def worksheet(self) -> WorksheetCollection:
        return WorksheetCollection(self)

    def close(self) -> None:
        if self._book is not None:
            self._book.close()
            self._book = None

    def __enter__(self) -> WorkbookSession:
        if self._book is None:
            raise _failure("workbook session is closed", ErrorCode.CLIENT_CLOSED)
        return self

    def __exit__(self, *_: object) -> None:
        self.close()

    def write(self, *, verify: bool = True) -> OperationResult[dict[str, Any]]:
        if self._book is None:
            raise _failure("workbook session is closed", ErrorCode.CLIENT_CLOSED)
        _, destination = _path_from_uri(self.uri)
        destination.parent.mkdir(parents=True, exist_ok=True)
        fd, temporary_name = tempfile.mkstemp(prefix=f".{destination.stem}.", suffix=".xlsx", dir=destination.parent)
        os.close(fd)
        temporary = Path(temporary_name)
        try:
            self._validate_limits()
            self._book.save(temporary)
            if self._opened and self.profile != _PROFILE:
                staged = _verify_xlsx(temporary, None, self.profile, self.limits, None)
                os.replace(temporary, destination)
                verified = staged.to_wire()
            elif destination.exists():
                raise _failure("destination already exists; create is exclusive", ErrorCode.DESTINATION_EXISTS)
            else:
                try:
                    os.link(temporary, destination)
                except FileExistsError:
                    raise _failure("destination already exists; create is exclusive", ErrorCode.DESTINATION_EXISTS) from None
                verified = self.verify().require_value() if verify else {"status": "written"}
            self._written = True
            return _result(verified, _WORKBOOK_WRITE, self.uri, operation="workbook.write", details=verified)
        except OTCError:
            if destination.exists() and not self._written and not self._opened:
                destination.unlink(missing_ok=True)
            raise
        except (OSError, ValueError) as exc:
            if not self._opened:
                destination.unlink(missing_ok=True)
            raise _failure("workbook could not be written", ErrorCode.EXECUTION_FAILED, reason=type(exc).__name__) from None
        finally:
            temporary.unlink(missing_ok=True)

    def verify(self) -> OperationResult[dict[str, Any]]:
        _, path = _path_from_uri(self.uri)
        if not path.is_file():
            raise _failure("workbook file does not exist", ErrorCode.TARGET_NOT_FOUND, path=str(path))
        try:
            verification = _verify_xlsx(path, self._expected or None, self.profile, self.limits, self._images or None)
        except OTCError:
            raise
        except Exception as exc:
            raise _failure("workbook verification failed", ErrorCode.ARTIFACT_INTEGRITY, reason=type(exc).__name__) from None
        return _result(verification.to_wire(), _WORKBOOK_VERIFY, self.uri, operation="workbook.verify", details=verification.to_wire())

    def inspect(self) -> OperationResult[dict[str, Any]]:
        self._validate_limits()
        value = {"profile": self.profile, "sheets": list(self._book.sheetnames), "written": self._written}
        return _result(value, CapabilityIdentity("spreadsheet.workbook.inspect", "1.0"), self.uri, operation="workbook.inspect", details=value)

    def _capture_expected(self) -> None:
        self._expected = {}
        for sheet in self._book.worksheets:
            cells: dict[str, dict[str, Any]] = {}
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cells[cell.coordinate] = {"type": "f" if cell.data_type == "f" else "s", "value": cell.value}
            self._expected[sheet.title] = cells

    def _validate_limits(self) -> None:
        if len(self._book.worksheets) > self.limits.sheets:
            raise _failure("workbook exceeds sheet limit", ErrorCode.RESOURCE_LIMIT)
        cells = 0
        text_bytes = 0
        for sheet in self._book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cells += 1
                        text_bytes += len(str(cell.value).encode())
        if cells > self.limits.cells or text_bytes > self.limits.text_bytes:
            raise _failure("workbook exceeds cell or text resource limits", ErrorCode.RESOURCE_LIMIT)
        if len(self._images) > self.limits.images:
            raise _failure("workbook exceeds image limit", ErrorCode.RESOURCE_LIMIT)


class WorksheetCollection:
    def __init__(self, book: WorkbookSession) -> None:
        self._book = book

    def create(self, name: str) -> Worksheet:
        if not isinstance(name, str) or not name.strip() or len(name) > 31 or any(char in name for char in "[]:*?/\\"):
            raise _failure("worksheet name is invalid", ErrorCode.INVALID_TARGET)
        if name in self._book._book.sheetnames:
            raise _failure("worksheet already exists", ErrorCode.KEY_CONFLICT, worksheet=name)
        if len(self._book._book.worksheets) >= self._book.limits.sheets:
            raise _failure("workbook exceeds sheet limit", ErrorCode.RESOURCE_LIMIT)
        sheet = self._book._book.create_sheet(name)
        self._book._expected.setdefault(name, {})
        return Worksheet(self._book, sheet)

    def __call__(self, name: str) -> Worksheet:
        return self.get(name)

    def get(self, name: str) -> Worksheet:
        try:
            return Worksheet(self._book, self._book._book[name])
        except KeyError:
            raise _failure("worksheet does not exist", ErrorCode.TARGET_NOT_FOUND, worksheet=name) from None

    def list(self) -> tuple[str, ...]:
        return tuple(self._book._book.sheetnames)


class Worksheet:
    def __init__(self, book: WorkbookSession, sheet: Any) -> None:
        self._book, self._sheet = book, sheet

    @property
    def name(self) -> str:
        return self._sheet.title

    def range(self, address: str) -> Range:
        try:
            normalized = RangeRef(address).address
        except ValueError as exc:
            raise _failure(str(exc), ErrorCode.INVALID_TARGET) from None
        return Range(self._book, self._sheet, normalized)

    def formulas(self) -> FormulaCollection:
        return FormulaCollection(self._book, self._sheet)

    def merge(self, address: str) -> OperationResult[None]:
        self.range(address)
        self._sheet.merge_cells(address.upper())
        return _result(None, _RANGE_WRITE, self._book.uri, operation="worksheet.merge", details={"worksheet": self.name, "range": address.upper()})

    def image(self, image: ImageSpec) -> OperationResult[None]:
        if not isinstance(image, ImageSpec):
            raise TypeError("image must be an ImageSpec")
        if len(image.content) > self._book.limits.image_bytes:
            raise _failure("image exceeds the per-image byte limit", ErrorCode.RESOURCE_LIMIT)
        if sum(item["byte_count"] for item in self._book._images) + len(image.content) > self._book.limits.total_image_bytes:
            raise _failure("workbook exceeds the total image byte limit", ErrorCode.RESOURCE_LIMIT)
        try:
            from openpyxl.drawing.image import Image as OpenpyxlImage

            picture = OpenpyxlImage(BytesIO(image.content))
            picture.anchor = image.anchor
            self._sheet.add_image(picture)
        except Exception as exc:
            raise _failure("image could not be added to the worksheet", ErrorCode.EXECUTION_FAILED, reason=type(exc).__name__) from None
        self._book._images.append({"sheet": self.name, "anchor": image.anchor, "sha256": image.sha256, "byte_count": len(image.content)})
        return _result(None, _IMAGE_INSERT, self._book.uri, operation="worksheet.image", details={"worksheet": self.name, "anchor": image.anchor, "sha256": image.sha256})

    def delete(self) -> OperationResult[None]:
        if len(self._book._book.worksheets) <= 1:
            raise _failure("a workbook must retain at least one worksheet", ErrorCode.INVALID_TARGET)
        self._book._book.remove(self._sheet)
        self._book._expected.pop(self.name, None)
        return _result(None, _RANGE_WRITE, self._book.uri, operation="worksheet.delete", details={"worksheet": self.name})


class Range:
    def __init__(self, book: WorkbookSession, sheet: Any, address: str) -> None:
        self._book, self._sheet, self.address = book, sheet, address

    def read(self) -> OperationResult[list[list[Any]]]:
        min_row, min_col, max_row, max_col = _coordinates(self.address)
        values = [[self._sheet.cell(row, col).value for col in range(min_col, max_col + 1)] for row in range(min_row, max_row + 1)]
        return _result(values, _RANGE_READ, self._book.uri, operation="range.read", details={"worksheet": self._sheet.title, "range": self.address})

    def write(self, values: Any, *, format: CellFormat | None = None, style: CellStyle | None = None) -> OperationResult[None]:
        min_row, min_col, max_row, max_col = _coordinates(self.address)
        if isinstance(values, (str, bytes)) or not isinstance(values, Sequence):
            matrix = [[values]]
        else:
            matrix = [list(row) if isinstance(row, Sequence) and not isinstance(row, (str, bytes)) else [row] for row in values]
        height, width = len(matrix), max((len(row) for row in matrix), default=0)
        if height != max_row - min_row + 1 or width != max_col - min_col + 1:
            raise _failure("range write dimensions do not match the selected range", ErrorCode.INVALID_TARGET)
        expected = self._book._expected.setdefault(self._sheet.title, {})
        for row_offset, row in enumerate(matrix):
            for col_offset, value in enumerate(row):
                cell = self._sheet.cell(min_row + row_offset, min_col + col_offset)
                text = _literal(value)
                cell.value = text
                cell.data_type = "s"
                cell.number_format = "@"
                if format is not None:
                    cell.number_format = format.pattern or format.kind
                _apply_style(cell, style)
                expected[cell.coordinate] = {"type": "s", "value": text}
        return _result(None, _RANGE_WRITE, self._book.uri, operation="range.write", details={"worksheet": self._sheet.title, "range": self.address})

    def style(self, value: CellStyle) -> OperationResult[None]:
        if not isinstance(value, CellStyle):
            raise TypeError("style must be a CellStyle")
        min_row, min_col, max_row, max_col = _coordinates(self.address)
        for row in self._sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                _apply_style(cell, value)
        return _result(None, _RANGE_STYLE, self._book.uri, operation="range.style", details={"worksheet": self._sheet.title, "range": self.address})

    def format(self, value: CellFormat) -> OperationResult[None]:
        if not isinstance(value, CellFormat):
            raise TypeError("format must be a CellFormat")
        min_row, min_col, max_row, max_col = _coordinates(self.address)
        pattern = value.pattern or value.kind
        for row in self._sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                cell.number_format = pattern
        return _result(None, _RANGE_FORMAT, self._book.uri, operation="range.format", details={"worksheet": self._sheet.title, "range": self.address, "format": pattern})

    def clear(self) -> OperationResult[None]:
        min_row, min_col, max_row, max_col = _coordinates(self.address)
        expected = self._book._expected.setdefault(self._sheet.title, {})
        for row in self._sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                cell.value = None
                expected.pop(cell.coordinate, None)
        return _result(None, _RANGE_WRITE, self._book.uri, operation="range.clear", details={"worksheet": self._sheet.title, "range": self.address})

    def sort(self, *, key_column: int = 1, reverse: bool = False) -> OperationResult[None]:
        min_row, min_col, max_row, max_col = _coordinates(self.address)
        if key_column < 1 or key_column > max_col - min_col + 1:
            raise _failure("sort key column is outside the selected range", ErrorCode.INVALID_TARGET)
        rows = [[self._sheet.cell(row, col).value for col in range(min_col, max_col + 1)] for row in range(min_row, max_row + 1)]
        rows.sort(key=lambda row: (row[key_column - 1] is None, row[key_column - 1]), reverse=reverse)
        for row_offset, values in enumerate(rows):
            for col_offset, value in enumerate(values):
                self._sheet.cell(min_row + row_offset, min_col + col_offset).value = value
        return _result(None, _RANGE_SORT, self._book.uri, operation="range.sort", details={"worksheet": self._sheet.title, "range": self.address})


class FormulaCollection:
    def __init__(self, book: WorkbookSession, sheet: Any) -> None:
        self._book, self._sheet = book, sheet

    def set(self, address: str, expression: str) -> OperationResult[FormulaWrite]:
        if not isinstance(expression, str) or not expression.startswith("="):
            raise _failure("formula expression must be a string beginning with '='", ErrorCode.INVALID_FORMULA)
        try:
            normalized = RangeRef(address).address
        except ValueError as exc:
            raise _failure(str(exc), ErrorCode.INVALID_TARGET) from None
        if ":" in normalized:
            raise _failure("formula set currently accepts one cell address", ErrorCode.INVALID_TARGET)
        if len(expression.encode()) > 8 * 1024:
            raise _failure("formula expression exceeds the target byte limit", ErrorCode.INVALID_FORMULA)
        self._sheet[normalized] = expression
        self._sheet[normalized].data_type = "f"
        self._book._expected.setdefault(self._sheet.title, {})[normalized] = {"type": "f", "value": expression}
        return _result(FormulaWrite(normalized, expression), _RANGE_WRITE, self._book.uri, operation="formula.set", details={"worksheet": self._sheet.title, "range": normalized})


def _apply_style(cell: Any, style: CellStyle | None) -> None:
    if style is None:
        return
    from openpyxl.styles import PatternFill

    font = copy(cell.font)
    if style.font_size is not None:
        font.sz = style.font_size
    if style.bold is not None:
        font.bold = style.bold
    if style.italic is not None:
        font.italic = style.italic
    if style.foreground:
        font.color = style.foreground[1:]
    cell.font = font
    if style.fill:
        cell.fill = PatternFill("solid", fgColor=style.fill[1:])


def _verify_xlsx(path: Path, expected: Mapping[str, Mapping[str, Mapping[str, Any]]] | None, profile: str, limits: ArtifactLimits, expected_images: Iterable[Mapping[str, Any]] | None = None) -> WorkbookVerification:
    try:
        with ZipFile(path) as archive:
            members = archive.namelist()
            if len(members) != len(set(members)):
                raise _failure("workbook archive contains duplicate members", ErrorCode.ARTIFACT_INTEGRITY)
            if len(members) > limits.zip_members:
                raise _failure("workbook archive exceeds member limit", ErrorCode.RESOURCE_LIMIT)
            if sum(info.file_size for info in archive.infolist()) > limits.decompressed_bytes:
                raise _failure("workbook archive exceeds decompressed byte limit", ErrorCode.RESOURCE_LIMIT)
    except BadZipFile:
        raise _failure("workbook is not a valid ZIP archive", ErrorCode.ARTIFACT_INTEGRITY) from None
    _verify_raw_xlsx(path, expected, profile)
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=False)
    try:
        if len(book.sheetnames) > limits.sheets:
            raise _failure("workbook exceeds sheet limit", ErrorCode.RESOURCE_LIMIT)
        formulas = 0
        cells = 0
        decoded: dict[str, dict[str, dict[str, Any]]] = {}
        for sheet in book.worksheets:
            values: dict[str, dict[str, Any]] = {}
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    cells += 1
                    if cell.data_type == "f":
                        formulas += 1
                        if profile == _PROFILE:
                            raise _failure("literal artifact contains a formula", ErrorCode.ARTIFACT_INTEGRITY, cell=cell.coordinate)
                    values[cell.coordinate] = {"type": "f" if cell.data_type == "f" else "s", "value": cell.value}
            decoded[sheet.title] = values
        if expected is not None and profile == _PROFILE:
            if tuple(decoded) != tuple(expected):
                raise _failure("workbook sheet order differs from the staged session", ErrorCode.READBACK_MISMATCH)
            if json.dumps(decoded, sort_keys=True, default=str) != json.dumps(expected, sort_keys=True, default=str):
                raise _failure("workbook cell coverage differs from the staged session", ErrorCode.READBACK_MISMATCH)
        actual_images: list[dict[str, Any]] = []
        for sheet in book.worksheets:
            for picture in sheet._images:
                anchor = picture.anchor
                coordinate = getattr(anchor, "_from", None)
                if coordinate is None or coordinate.colOff or coordinate.rowOff:
                    raise _failure("image anchor must be a cell-aligned one-cell anchor", ErrorCode.ARTIFACT_INTEGRITY)
                from openpyxl.utils import get_column_letter

                actual_images.append({"sheet": sheet.title, "anchor": f"{get_column_letter(coordinate.col + 1)}{coordinate.row + 1}", "sha256": hashlib.sha256(picture._data()).hexdigest()})
        if expected_images is not None:
            wanted = [{"sheet": item["sheet"], "anchor": item["anchor"], "sha256": item["sha256"]} for item in expected_images]
            if actual_images != wanted:
                raise _failure("workbook image coverage or hashes differ", ErrorCode.READBACK_MISMATCH)
        semantic = hashlib.sha256(json.dumps({"sheets": decoded, "images": actual_images}, sort_keys=True, default=str, separators=(",", ":")).encode()).hexdigest()
        content = hashlib.sha256(path.read_bytes()).hexdigest()
        return WorkbookVerification("verified", f"sha256:{content}", f"sha256:{semantic}", tuple(book.sheetnames), cells, formulas, len(actual_images))
    finally:
        book.close()


def _verify_raw_xlsx(path: Path, expected: Mapping[str, Mapping[str, Mapping[str, Any]]] | None, profile: str) -> None:
    """Validate worksheet XML before openpyxl can normalize malformed cells."""

    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    with ZipFile(path) as archive:
        try:
            workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        except (KeyError, ElementTree.ParseError):
            raise _failure("workbook XML parts are incomplete", ErrorCode.ARTIFACT_INTEGRITY) from None
        relation_by_id: dict[str, Any] = {}
        for relation in relationships:
            identity = relation.get("Id")
            if not identity or identity in relation_by_id:
                raise _failure("workbook relationship identity is duplicated", ErrorCode.ARTIFACT_INTEGRITY)
            relation_by_id[identity] = relation
        sheets = workbook.findall(f"{ns}sheets/{ns}sheet")
        names = [sheet.get("name") for sheet in sheets]
        if any(not isinstance(name, str) or not name for name in names):
            raise _failure("workbook contains an unnamed worksheet", ErrorCode.ARTIFACT_INTEGRITY)
        if expected is not None and names != list(expected):
            raise _failure("raw worksheet sheet order differs", ErrorCode.READBACK_MISMATCH)
        for sheet in sheets:
            relation = relation_by_id.get(sheet.get(rel_ns + "id"))
            if relation is None or relation.get("TargetMode") == "External" or not relation.get("Type", "").endswith("/worksheet"):
                raise _failure("worksheet relationship is invalid", ErrorCode.ARTIFACT_INTEGRITY)
            target = relation.get("Target", "")
            member = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target))
            try:
                root = ElementTree.fromstring(archive.read(member))
            except (KeyError, ElementTree.ParseError):
                raise _failure("worksheet XML part is incomplete", ErrorCode.ARTIFACT_INTEGRITY) from None
            coordinates: set[str] = set()
            semantic: set[str] = set()
            for cell in root.iter(ns + "c"):
                coordinate = cell.get("r", "")
                if re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", coordinate) is None or coordinate in coordinates:
                    raise _failure("worksheet contains an invalid or duplicate cell coordinate", ErrorCode.ARTIFACT_INTEGRITY, cell=coordinate)
                coordinates.add(coordinate)
                if cell.find(ns + "f") is not None:
                    if profile == _PROFILE:
                        raise _failure("literal artifact contains a formula", ErrorCode.ARTIFACT_INTEGRITY, cell=coordinate)
                    semantic.add(coordinate)
                elif cell.get("t", "n") != "n" or cell.find(ns + "v") is not None or cell.find(ns + "is") is not None:
                    semantic.add(coordinate)
            if expected is not None and profile == _PROFILE:
                sheet_name = sheet.get("name")
                if semantic != set(expected.get(sheet_name, {})):
                    raise _failure("raw worksheet semantic coordinate coverage differs", ErrorCode.READBACK_MISMATCH, worksheet=sheet_name)


__all__ = ["FormulaWrite", "Range", "WorkbookSession", "WorkbookVerification", "Worksheet", "WorksheetCollection"]
