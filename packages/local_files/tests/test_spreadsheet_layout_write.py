from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from open_table_connector.local_files.spreadsheet_workbook import LocalSpreadsheetProvider
from open_table_connector.spreadsheets import Change, SpreadsheetTarget
from open_table_connector.contract import CapabilityIdentity


def _change(operation, arguments, target="Report"):
    return Change(
        operation,
        CapabilityIdentity("spreadsheet." + operation, "1.0"),
        target,
        arguments,
    )


def test_layout_writer_preserves_values_and_applies_financial_style(tmp_path: Path):
    path = tmp_path / "layout.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Report"
    sheet["A1"] = 60
    sheet["B1"] = "amount"
    book.save(path)
    provider = LocalSpreadsheetProvider()
    binding = provider.bind(SpreadsheetTarget(path.as_uri()))
    binding["revision"] = None
    changes = [
        _change("range.format", {"address": "A1", "pattern": '[$€-407] #,##0.00'}),
        _change("range.style", {"address": "B1", "style": {
            "bold": True,
            "horizontal": "center",
            "vertical": "center",
            "wrap_text": True,
            "shrink_to_fit": False,
            "border": {"top": {"style": "thin", "color": "#000000"}},
        }}),
        _change("worksheet.config", {"row_heights": {"1": 24}, "column_sizes": {"A": {"value": 18, "unit": "excel_character"}}, "show_gridlines": False}),
    ]
    provider.commit(binding, changes)
    saved = openpyxl.load_workbook(path, data_only=False)
    sheet = saved["Report"]
    assert sheet["A1"].value == 60
    assert sheet["A1"].number_format == '[$€-407] #,##0.00'
    assert sheet["B1"].font.bold is True
    assert sheet["B1"].alignment.horizontal == "center"
    assert sheet["B1"].alignment.wrap_text is True
    assert sheet["B1"].alignment.shrink_to_fit in (False, None)
    assert sheet["B1"].border.top.style == "thin"
    assert sheet.row_dimensions[1].height == 24
    assert sheet.column_dimensions["A"].width == 18
    assert sheet.sheet_view.showGridLines is False


def test_layout_writer_rejects_clip_without_mutating(tmp_path: Path):
    path = tmp_path / "layout.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Report"
    book.save(path)
    provider = LocalSpreadsheetProvider()
    binding = provider.bind(SpreadsheetTarget(path.as_uri()))
    with pytest.raises(Exception):
        provider.preflight(binding, [_change("range.style", {"address": "A1", "style": {"text_layout": "clip"}})])
