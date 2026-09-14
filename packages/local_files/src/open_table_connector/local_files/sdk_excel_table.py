"""Portable, create-only Excel worksheet materialization."""

from __future__ import annotations

import hashlib
import json
import uuid
from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from urllib.parse import parse_qsl, unquote, urlsplit

import polars as pl
from open_table_connector.contract import ConnectorError
from open_table_connector.sdk._excel_table import table_address, table_matrix
from open_table_connector.sdk.materialization import MaterializationRequest
from open_table_connector.sdk.model import (
    DirectDestination,
    SheetModeDestination,
    SheetModeTableAddress,
    _schema_from_wire,
    _schema_to_wire,
)
from open_table_connector.sdk.result import (
    CommitState,
    ErrorCode,
    ErrorInfo,
    OperationResult,
    OperationWarning,
    OTCError,
    Outcome,
    Receipt,
    VerificationState,
)
from open_table_connector.sdk.table import TableBinding
from open_table_connector.spreadsheets import ArtifactLimits
from open_table_connector.spreadsheets._session import SpreadsheetSession
from open_table_connector.spreadsheets.model import SpreadsheetTarget

from .spreadsheet_workbook import LocalSpreadsheetProvider

_METADATA_PREFIX = "_otc_materialization_"
_METADATA_MAGIC = "otc.portable-excel-metadata"


def _column_name(index):
    name = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name


def _workbook_evidence(written, connector, uri):
    receipts = tuple(
        Receipt(
            "workbook-provider",
            item.get("operation", "workbook.write"),
            connector.identity.connector_id,
            safe_target=uri,
            mode="sheet-mode",
            details=item.get("details", {}),
        )
        for item in written.get("receipts", ())
    )
    warnings = tuple(
        item
        if isinstance(item, OperationWarning)
        else OperationWarning(item["code"], item["message"], item.get("safe_details", {}))
        for item in written.get("warnings", ())
    )
    return receipts, warnings


def _destination(destination):
    if isinstance(destination, SheetModeDestination):
        uri, sheet, anchor, header = (
            destination.grid.value,
            destination.worksheet,
            destination.anchor,
            destination.header,
        )
    elif isinstance(destination, DirectDestination):
        parsed = urlsplit(destination.uri.value)
        selectors = parse_qsl(parsed.fragment, keep_blank_values=True)
        if len(selectors) != 1 or selectors[0][0] != "sheet" or not selectors[0][1]:
            raise ValueError("Excel materialization requires file:///absolute/book.xlsx#sheet=Name")
        uri, sheet, anchor, header = (
            parsed._replace(fragment="").geturl(),
            selectors[0][1],
            "A1",
            True,
        )
    else:
        raise ValueError("Excel materialization requires a sheet destination")
    parsed = urlsplit(uri)
    if (
        parsed.scheme != "file"
        or parsed.netloc not in ("", "localhost")
        or parsed.query
        or parsed.fragment
    ):
        raise ValueError("Excel materialization requires an absolute file URI")
    path = Path(unquote(parsed.path))
    if (
        not path.is_absolute()
        or path.suffix.lower() != ".xlsx"
        or not sheet
        or anchor != "A1"
        or header is not True
    ):
        raise ValueError("Excel portable sheet mode requires worksheet, anchor=A1, and header=true")
    return path, sheet, path.as_uri()


def _cell(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _metadata_name(table_id, names):
    base = _METADATA_PREFIX + table_id.replace("-", "")[:10]
    candidate, number = base, 1
    folded = {name.casefold() for name in names}
    while candidate.casefold() in folded:
        number += 1
        candidate = f"{base}_{number}"
    return candidate


def _metadata(source, table_id, worksheet, grid):
    return json.dumps(
        {
            "magic": _METADATA_MAGIC,
            "kind": "sheet-mode-table",
            "version": 1,
            "grid": grid,
            "table_id": table_id,
            "worksheet": worksheet,
            "anchor": "A1",
            "header": True,
            "schema": _schema_to_wire(source.schema),
            "rows": source.height,
            "schema_fingerprint": MaterializationRequest(
                source,
                DirectDestination("file:///placeholder.xlsx"),
                "otc.portable-table/v1",
                "metadata",
            ).schema_fingerprint,
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )


def _find_metadata(path, table_id=None, worksheet=None):
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=False)
    try:
        records = []
        for sheet in book.worksheets:
            if not sheet.title.startswith(_METADATA_PREFIX) or sheet.sheet_state != "veryHidden":
                continue
            raw = sheet["A1"].value
            try:
                record = json.loads(raw) if isinstance(raw, str) else None
            except json.JSONDecodeError:
                continue
            if not isinstance(record, dict) or set(record) != {
                "magic",
                "kind",
                "version",
                "grid",
                "table_id",
                "worksheet",
                "anchor",
                "header",
                "schema",
                "rows",
                "schema_fingerprint",
            }:
                continue
            if (
                record["magic"] != _METADATA_MAGIC
                or record["kind"] != "sheet-mode-table"
                or record["version"] != 1
                or record["grid"] != path.as_uri()
                or record["anchor"] != "A1"
                or record["header"] is not True
            ):
                continue
            if (table_id is None or record["table_id"] == table_id) and (
                worksheet is None or record["worksheet"] == worksheet
            ):
                records.append(record)
        if len(records) != 1:
            raise ValueError("portable Excel schema metadata is missing, stale, or malformed")
        record = records[0]
        if record["worksheet"] not in book.sheetnames or record["rows"] < 0:
            raise ValueError("portable Excel schema metadata is stale")
        schema = _schema_from_wire(record["schema"])
        if schema is None:
            raise ValueError("portable Excel schema metadata is malformed")
        expected = MaterializationRequest(
            pl.DataFrame(schema=schema),
            DirectDestination("file:///placeholder.xlsx"),
            "otc.portable-table/v1",
            "metadata",
        ).schema_fingerprint
        if record["schema_fingerprint"] != expected:
            raise ValueError("portable Excel schema metadata fingerprint is invalid")
        return record, schema
    finally:
        book.close()


def _typed_frame(path, record, schema):
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book[record["worksheet"]]
        headers = [sheet.cell(1, index + 1).value for index in range(len(schema))]
        if headers != list(schema.names()):
            raise ValueError("portable Excel worksheet does not match its metadata")
        rows = [
            [
                (
                    ""
                    if sheet.cell(row, column + 1).data_type == "inlineStr"
                    and sheet.cell(row, column + 1).value is None
                    else sheet.cell(row, column + 1).value
                )
                for column in range(len(schema))
            ]
            for row in range(2, record["rows"] + 2)
        ]
        if any(not isinstance(value, (str, type(None))) for row in rows for value in row):
            raise ValueError("portable Excel worksheet contains non-lexical values")
        values = {name: [row[index] for row in rows] for index, name in enumerate(schema.names())}
        frame = pl.DataFrame(values, schema={name: pl.String for name in schema.names()})
        expressions = []
        for name, dtype in schema.items():
            if dtype == pl.String:
                expressions.append(pl.col(name))
            elif dtype == pl.Boolean:
                if not frame.get_column(name).drop_nulls().is_in(["true", "false"]).all():
                    raise ValueError("portable Excel boolean is not canonical")
                expressions.append(
                    pl.when(pl.col(name).is_null())
                    .then(None)
                    .otherwise(pl.col(name) == "true")
                    .cast(dtype)
                    .alias(name)
                )
            else:
                expressions.append(pl.col(name).cast(dtype, strict=True).alias(name))
        return frame.select(expressions)
    finally:
        book.close()


def open_portable_excel(address):
    path = (
        Path(unquote(urlsplit(address.grid.value).path))
        if isinstance(address, SheetModeTableAddress)
        else None
    )
    table_id = address.table_id if isinstance(address, SheetModeTableAddress) else None
    worksheet = None
    if path is None:
        path, worksheet, _ = _destination(DirectDestination(address))
    record, schema = _find_metadata(path, table_id=table_id, worksheet=worksheet)
    frame = _typed_frame(path, record, schema)
    canonical = SheetModeTableAddress(path.as_uri(), record["table_id"])
    revision = "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()
    return TableBinding(
        path.as_uri(),
        "sheet-mode",
        frame.schema,
        revision,
        "local-files",
        "otc.portable-table/v1",
        frame.height,
        record["schema_fingerprint"],
        None,
        canonical,
    ), frame


def create_portable_excel_table(connector, request: MaterializationRequest):
    written = None
    mutation = ()
    commit_receipts = ()
    warnings = ()
    try:
        path, worksheet, uri = _destination(request.destination)
        table_id = str(uuid.uuid4())
        provider = LocalSpreadsheetProvider()
        session = SpreadsheetSession(
            provider, SpreadsheetTarget(uri), new=not path.exists(), profile="general/1.0"
        )
        try:
            existing = set(session.observe("worksheet.list")["value"])
            if worksheet.casefold() in {name.casefold() for name in existing}:
                return OperationResult(
                    None,
                    Outcome.REJECTED,
                    CommitState.NOT_STARTED,
                    VerificationState.SKIPPED,
                    (),
                    error=ErrorInfo(
                        ErrorCode.DESTINATION_EXISTS, "Excel destination worksheet already exists"
                    ),
                )
            metadata_sheet = _metadata_name(table_id, existing | {worksheet})
            session.queue("worksheet.create", worksheet, {"name": worksheet})
            session.queue(
                "range.write",
                worksheet,
                {
                    "address": f"A1:{_column_name(request.source.width)}{request.source.height + 1}",
                    "values": [
                        request.source.columns,
                        *[[_cell(value) for value in row] for row in request.source.iter_rows()],
                    ],
                },
            )
            session.queue("worksheet.create", metadata_sheet, {"name": metadata_sheet})
            session.queue(
                "range.write",
                metadata_sheet,
                {
                    "address": "A1",
                    "values": [[_metadata(request.source, table_id, worksheet, uri)]],
                },
            )
            session.queue(
                "worksheet.config", metadata_sheet, {"properties": {"visibility": "veryHidden"}}
            )
            written = session.write(expected_revision=session.binding.get("revision"))
        finally:
            session.close()
        mutation = (
            Receipt(
                "physical",
                "table.materialize.create",
                connector.identity.connector_id,
                "table.materialize.create/1.0",
                uri,
                "sheet-mode",
                {"revision": written["binding"]["revision"]},
            ),
        )
        commit_receipts, warnings = _workbook_evidence(written, connector, uri)
        address = SheetModeTableAddress(uri, table_id)
        binding, observed = open_portable_excel(address)
        read = Receipt(
            "physical",
            "table.read",
            connector.identity.connector_id,
            "table.read/1.0",
            uri,
            "sheet-mode",
            {"revision": binding.observed_revision},
        )
        if not observed.equals(request.source):
            return OperationResult(
                None,
                Outcome.FAILED,
                CommitState.COMMITTED,
                VerificationState.FAILED,
                (*commit_receipts, *mutation, read),
                warnings=warnings,
                error=ErrorInfo(
                    ErrorCode.READBACK_MISMATCH,
                    "portable Excel readback differs from submitted table",
                    {"expected": str(request.source), "observed": str(observed)},
                ),
            )
        return OperationResult(
            binding,
            Outcome.SUCCEEDED,
            CommitState.COMMITTED,
            VerificationState.PASSED,
            (*commit_receipts, *mutation, read),
            warnings=warnings,
        )
    except ConnectorError as exc:
        if written is not None:
            return OperationResult(
                None,
                Outcome.FAILED,
                CommitState.COMMITTED,
                VerificationState.FAILED,
                (*commit_receipts, *mutation),
                warnings=warnings,
                error=ErrorInfo(
                    ErrorCode.READBACK_MISMATCH,
                    "portable Excel commit completed but readback failed",
                    {"cause": type(exc).__name__},
                ),
            )
        code = (
            ErrorCode.STALE_REVISION
            if exc.code.value == "conflict" and "stale" in exc.message.lower()
            else ErrorCode.EXECUTION_FAILED
        )
        return OperationResult(
            None,
            Outcome.REJECTED,
            CommitState.NOT_STARTED,
            VerificationState.SKIPPED,
            (),
            error=ErrorInfo(code, exc.message, dict(exc.safe_details)),
        )
    except Exception as exc:
        if written is not None:
            return OperationResult(
                None,
                Outcome.FAILED,
                CommitState.COMMITTED,
                VerificationState.FAILED,
                (*commit_receipts, *mutation),
                warnings=warnings,
                error=ErrorInfo(
                    ErrorCode.READBACK_MISMATCH,
                    "portable Excel commit completed but readback failed",
                    {"cause": type(exc).__name__},
                ),
            )
        return OperationResult(
            None,
            Outcome.REJECTED,
            CommitState.NOT_STARTED,
            VerificationState.SKIPPED,
            (),
            error=ErrorInfo(ErrorCode.INVALID_TARGET, str(exc)),
        )


def create_excel_table(connector, source, destination):
    """Legacy lexical implementation remains intentionally separate."""
    from open_table_connector.sdk import Client, ConnectorRegistry

    client = book = written = None
    receipts = ()
    try:
        if not isinstance(destination, DirectDestination):
            raise ValueError("Excel materialization requires DirectDestination")
        uri = destination.uri
        parsed = urlsplit(uri.value)
        selectors = parse_qsl(parsed.fragment, keep_blank_values=True)
        if (
            parsed.scheme != "file"
            or parsed.netloc not in ("", "localhost")
            or parsed.query
            or len(selectors) != 1
            or selectors[0][0] != "sheet"
            or not selectors[0][1]
        ):
            raise ValueError("Excel materialization requires file:///absolute/book.xlsx#sheet=Name")
        path, sheet = Path(unquote(parsed.path)), selectors[0][1]
        if not path.is_absolute() or path.suffix.lower() != ".xlsx":
            raise ValueError("Excel materialization requires an absolute .xlsx file")
        if (source.height + 1) * source.width > ArtifactLimits().cells:
            raise ValueError("Excel table exceeds cell limit")
        matrix = table_matrix(source)
        for column, dtype in enumerate(source.dtypes):
            if dtype.is_integer():
                for index, value in enumerate(source.get_column(source.columns[column])):
                    if value is not None and len(str(abs(value))) <= 15:
                        matrix[index + 1][column] = value
        if any(all(value is None for value in row) for row in matrix[1:]):
            raise ValueError("worksheet table cannot preserve all-null rows")
        client = Client(registry=ConnectorRegistry([connector]))
        book = (
            client.workbook.open(path.as_uri())
            if path.exists()
            else client.workbook.create(path.as_uri(), profile="general/1.0")
        )
        if sheet.casefold() in {name.casefold() for name in book.worksheet.list()}:
            raise ValueError("Excel destination worksheet already exists")
        book.worksheet.create(sheet).range(table_address(len(matrix), source.width)).write(matrix)
        written = book.write()
        receipts = written.receipts
        if (
            written.outcome is not Outcome.SUCCEEDED
            or written.commit is not CommitState.COMMITTED
            or written.verification is not VerificationState.PASSED
        ):
            return replace(written, value=None)
        opened = connector.open_table(uri.value)
        receipts += opened.receipts
        binding = opened.require_value()
        readback = connector.read_table(binding)
        receipts += readback.receipts
        if not readback.require_value().to_polars().equals(source.cast(pl.String)):
            raise ValueError("persisted Excel table differs from submitted lexical values")
        return OperationResult(
            binding, Outcome.SUCCEEDED, CommitState.COMMITTED, VerificationState.PASSED, receipts
        )
    except Exception as exc:
        if written is not None:
            return OperationResult(
                None,
                Outcome.FAILED,
                written.commit,
                VerificationState.FAILED,
                receipts,
                error=ErrorInfo(
                    ErrorCode.EXECUTION_FAILED,
                    "Excel table readback failed",
                    {"cause": str(exc), "commit": written.commit.value},
                ),
            )
        if isinstance(exc, OTCError):
            return exc.result
        return OperationResult(
            None,
            Outcome.REJECTED,
            CommitState.NOT_STARTED,
            VerificationState.SKIPPED,
            (),
            error=ErrorInfo(ErrorCode.INVALID_TARGET, str(exc)),
        )
    finally:
        if book is not None:
            book.close()
        if client is not None:
            client.close()
