from datetime import UTC, date, datetime
from decimal import Decimal

import open_table_connector.sdk as otc
import polars as pl
import pytest
from open_table_connector.local_files import LocalFilesConnector
from open_table_connector.sdk import (
    PORTABLE_TABLE_PROFILE_V1,
    Client,
    ConnectorRegistry,
    ErrorCode,
    OTCError,
    SheetModeDestination,
)
from open_table_connector.sdk.model import SheetModeTableAddress
from openpyxl import Workbook, load_workbook


def _client():
    return Client(registry=ConnectorRegistry([LocalFilesConnector()]))


def _portable(client, frame, destination, *, key=None):
    return client.materialize(
        frame,
        to=destination,
        profile=PORTABLE_TABLE_PROFILE_V1,
        idempotency_key=key or f"excel-portable-test:{destination}",
    )


def test_portable_excel_uses_structured_destination_and_recovers_exact_types(tmp_path):
    path = tmp_path / "typed.xlsx"
    source = pl.DataFrame(
        {
            "text": ["東京", "", None],
            "flag": [True, False, None],
            "integer": [9223372036854775807, -9223372036854775808, None],
            "float": [1.5, None, -0.0],
            "decimal": [Decimal("123.40"), None, Decimal("-0.01")],
            "day": [date(2024, 1, 2), None, date(1999, 12, 31)],
            "stamp": [
                datetime(2024, 1, 2, 3, 4, 5, 678901, tzinfo=UTC),
                None,
                datetime(1999, 12, 31, 23, 59, 59, tzinfo=UTC),
            ],
        },
        schema={
            "text": pl.String,
            "flag": pl.Boolean,
            "integer": pl.Int64,
            "float": pl.Float64,
            "decimal": pl.Decimal(precision=10, scale=2),
            "day": pl.Date,
            "stamp": pl.Datetime("us", "UTC"),
        },
    )
    destination = SheetModeDestination(path.as_uri(), "A1", True, "Base")

    result = _portable(_client(), source, destination)

    address = result.require_value().address
    assert isinstance(address, SheetModeTableAddress)
    assert address.grid.value == path.as_uri()
    assert address.table_id
    assert SheetModeTableAddress.from_wire(address.to_wire()) == address
    assert _client().open(address).require_value().read().require_value().equals(source)


def test_portable_excel_all_null_and_zero_rows_require_durable_metadata(tmp_path):
    path = tmp_path / "empty.xlsx"
    source = pl.DataFrame(
        {"empty": [None, None]}, schema={"empty": pl.Decimal(precision=8, scale=3)}
    )
    result = _portable(_client(), source, f"file://{path}#sheet=Base")
    assert (
        _client()
        .open(result.require_value().address)
        .require_value()
        .read()
        .require_value()
        .equals(source)
    )

    zero = pl.DataFrame(schema={"id": pl.Int64, "when": pl.Datetime("us", "UTC")})
    _portable(_client(), zero, f"file://{path}#sheet=Zero")
    assert (
        _client()
        .open(f"file://{path}#sheet=Zero")
        .require_value()
        .read()
        .require_value()
        .equals(zero)
    )


def test_portable_excel_rejects_casefold_conflict_and_preserves_unrelated_sheet(tmp_path):
    path = tmp_path / "existing.xlsx"
    book = Workbook()
    report = book.active
    report.title = "Report"
    report["A1"] = "Title"
    report["A1"].font = report["A1"].font.copy(bold=True)
    report["B1"] = "=1+2"
    book.save(path)
    book.close()

    result = _portable(_client(), pl.DataFrame({"id": [1]}), f"file://{path}#sheet=Base")
    assert result.commit.value == "committed"
    before = path.read_bytes()
    with pytest.raises(OTCError) as caught:
        _portable(_client(), pl.DataFrame({"id": [2]}), f"file://{path}#sheet=base")
    assert caught.value.result.error.code is ErrorCode.DESTINATION_EXISTS
    assert path.read_bytes() == before
    observed = load_workbook(path, data_only=False)
    try:
        assert observed["Report"]["A1"].value == "Title"
        assert observed["Report"]["A1"].font.bold is True
        assert observed["Report"]["B1"].value == "=1+2"
    finally:
        observed.close()


def test_portable_excel_missing_or_tampered_metadata_never_infers_schema(tmp_path):
    path = tmp_path / "tampered.xlsx"
    result = _portable(_client(), pl.DataFrame({"id": [1]}), f"file://{path}#sheet=Base")
    address = result.require_value().address
    book = load_workbook(path)
    try:
        metadata = next(
            sheet for sheet in book.worksheets if sheet.title.startswith("_otc_materialization_")
        )
        metadata["A1"] = "tampered"
        book.save(path)
    finally:
        book.close()
    with pytest.raises(OTCError) as caught:
        _client().open(address)
    assert caught.value.result.error.code is ErrorCode.INVALID_SCHEMA


def test_portable_excel_uses_excel_columns_beyond_z_and_rejects_bad_boolean(tmp_path):
    path = tmp_path / "wide.xlsx"
    source = pl.DataFrame({**{f"c{index}": [index] for index in range(28)}, "flag": [True]})
    result = _portable(_client(), source, f"file://{path}#sheet=Wide")
    assert result.require_value().read().require_value().equals(source)
    address = result.require_value().address
    book = load_workbook(path)
    try:
        book["Wide"]["AC2"] = "truthy"
        book.save(path)
    finally:
        book.close()
    with pytest.raises(OTCError) as caught:
        _client().open(address)
    assert caught.value.result.error.code is ErrorCode.INVALID_SCHEMA


def test_portable_excel_committed_mismatch_receipts_are_flat_and_ordered(tmp_path, monkeypatch):
    import open_table_connector.local_files.sdk_excel_table as excel

    original = excel.open_portable_excel

    def mismatch(address):
        binding, frame = original(address)
        return binding, frame.head(0)

    monkeypatch.setattr(excel, "open_portable_excel", mismatch)
    with pytest.raises(OTCError) as caught:
        _portable(
            _client(), pl.DataFrame({"id": [1]}), f"file://{tmp_path / 'receipt.xlsx'}#sheet=Base"
        )
    result = caught.value.result
    assert result.error.code is ErrorCode.READBACK_MISMATCH
    assert [receipt.operation for receipt in result.receipts] == [
        "workbook.write",
        "table.materialize.create",
        "table.read",
    ]


def test_portable_excel_preserves_workbook_commit_evidence(tmp_path, monkeypatch):
    from open_table_connector.spreadsheets._session import SpreadsheetSession

    original = SpreadsheetSession.write

    def write_with_evidence(self, **kwargs):
        result = dict(original(self, **kwargs))
        result["receipts"] = (
            {"operation": "workbook.provider.commit", "details": {"injected": True}},
        )
        result["warnings"] = ({"code": "injected", "message": "provider warning"},)
        return result

    monkeypatch.setattr(SpreadsheetSession, "write", write_with_evidence)
    result = _portable(
        _client(), pl.DataFrame({"id": [1]}), f"file://{tmp_path / 'evidence.xlsx'}#sheet=Base"
    )
    assert [receipt.operation for receipt in result.receipts] == [
        "workbook.provider.commit",
        "table.materialize.create",
        "table.read",
    ]
    assert [warning.code for warning in result.warnings] == ["injected"]


def test_portable_excel_replays_by_key_and_redacts_provider_snapshots(tmp_path):
    source = pl.DataFrame({"secret": ["unique-submit-secret"]})
    destination = f"file://{tmp_path / 'replay.xlsx'}#sheet=Base"
    first = _portable(_client(), source, destination)
    replay = _portable(_client(), source, destination)
    assert replay.require_value().address == first.require_value().address
    assert "unique-submit-secret" not in repr(tuple(receipt.to_wire() for receipt in replay.receipts))


def test_portable_excel_scopes_reused_key_to_the_canonical_worksheet(tmp_path):
    source = pl.DataFrame({"id": [1]})
    path = tmp_path / "scoped.xlsx"
    first = _portable(_client(), source, f"file://{path}#sheet=First", key="same-key")
    second = _portable(_client(), source, f"file://{path}#sheet=Second", key="same-key")
    assert second.require_value().address.table_id != first.require_value().address.table_id


def test_portable_excel_maps_stable_provider_conflict_to_stale_revision(tmp_path, monkeypatch):
    from open_table_connector.contract import ConnectorError, ConnectorErrorCode
    from open_table_connector.spreadsheets._session import SpreadsheetSession

    def conflict(self, **kwargs):
        del self, kwargs
        raise ConnectorError(ConnectorErrorCode.CONFLICT, "revision changed", {"reason": "revision_changed"})

    monkeypatch.setattr(SpreadsheetSession, "write", conflict)
    with pytest.raises(OTCError) as raised:
        _portable(_client(), pl.DataFrame({"id": [1]}), f"file://{tmp_path / 'stale.xlsx'}#sheet=Base")
    assert raised.value.result.error.code is ErrorCode.STALE_REVISION
    assert raised.value.result.commit is otc.CommitState.NOT_COMMITTED
