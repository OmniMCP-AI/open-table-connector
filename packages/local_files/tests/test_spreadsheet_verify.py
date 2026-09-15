from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from open_table_connector.contract import ConnectorError
from open_table_connector.local_files.spreadsheet_verify import verify_snapshot
from openpyxl import Workbook


def fixture():
    book = Workbook()
    book.active.title = "Report"
    book.active["A1"] = "=literal"
    book.active["A1"].data_type = "s"
    book.active["A1"].number_format = "@"
    book.calculation = None
    out = BytesIO()
    book.save(out)
    return out.getvalue(), {
        "schema": "otc.xlsx-physical/1.0",
        "profile": "literal-artifact/1.0",
        "sheets": [
            {
                "name": "Report",
                "cells": {"A1": {"value": "=literal", "type": "s", "format": "@", "style": {}}},
                "config": {},
                "merges": [],
                "images": [],
            }
        ],
    }


def corrupt(data, part, replace):
    out = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(out, "w", ZIP_DEFLATED) as target:
        for name in source.namelist():
            value = source.read(name)
            target.writestr(name, replace(value) if name == part else value)
    return out.getvalue()


def test_independent_fixture():
    data, expected = fixture()
    assert verify_snapshot(data, expected)["cells"] == 1


@pytest.mark.parametrize(
    "part,replace",
    [
        ("xl/workbook.xml", lambda b: b.replace(b"</workbook>", b"<calcPr/></workbook>")),
        (
            "xl/worksheets/sheet1.xml",
            lambda b: b.replace(
                b"</sheetData>",
                b'<row r="2"><c r="A2" t="inlineStr"><is><t>extra</t></is></c></row></sheetData>',
            ),
        ),
        (
            "xl/worksheets/sheet1.xml",
            lambda b: b.replace(b"</sheetData>", b'<row r="2"><c r="A1"/></row></sheetData>'),
        ),
        ("xl/worksheets/sheet1.xml", lambda b: b.replace(b'r="A1"', b'r="XFE1"')),
        ("xl/worksheets/sheet1.xml", lambda b: b.replace(b"<is>", b"<f>1+1</f><is>")),
        ("xl/worksheets/sheet1.xml", lambda b: b'<!DOCTYPE x [<!ENTITY x "bad">]>' + b),
        (
            "xl/_rels/workbook.xml.rels",
            lambda b: b.replace(b'Target="/xl/worksheets/sheet1.xml"', b'Target="../../evil.xml"'),
        ),
    ],
)
def test_corruption_fails_closed(part, replace):
    data, expected = fixture()
    with pytest.raises(ConnectorError):
        verify_snapshot(corrupt(data, part, replace), expected)


def test_limit_and_expectation_mismatch():
    data, expected = fixture()
    with pytest.raises(ConnectorError):
        verify_snapshot(data, expected, {"decompressed_bytes": 10})
    expected["sheets"][0]["cells"]["A1"]["value"] = "other"
    with pytest.raises(ConnectorError):
        verify_snapshot(data, expected)


def test_duplicate_member_and_orphan_part():
    data, expected = fixture()
    for name in ("xl/workbook.xml", "xl/media/image99.png", "xl/vbaProject.bin"):
        out = BytesIO(data)
        with ZipFile(out, "a") as archive:
            archive.writestr(name, b"bad")
        with pytest.raises(ConnectorError):
            verify_snapshot(out.getvalue(), expected)


def test_style_and_native_layout_tampering():
    data, expected = fixture()
    expected["sheets"][0]["cells"]["A1"]["style"] = {"bold": True}
    with pytest.raises(ConnectorError):
        verify_snapshot(data, expected)
    expected["sheets"][0]["cells"]["A1"]["style"] = {}
    expected["sheets"][0]["config"] = {"orientation": "landscape"}
    with pytest.raises(ConnectorError):
        verify_snapshot(data, expected)


def test_empty_literal_coverage():
    data, expected = fixture()
    data = corrupt(data, "xl/worksheets/sheet1.xml", lambda b: b.replace(b"=literal", b""))
    expected["sheets"][0]["cells"]["A1"]["value"] = ""
    assert verify_snapshot(data, expected)["cells"] == 1


def test_general_archive_enforces_xml_and_expansion():
    from open_table_connector.local_files.spreadsheet_verify import read_archive

    data, _ = fixture()
    assert "xl/workbook.xml" in read_archive(data)
    with pytest.raises(ConnectorError):
        read_archive(data, {"member_bytes": 1})
    data = corrupt(data, "xl/workbook.xml", lambda b: b"<!DOCTYPE x>" + b)
    with pytest.raises(ConnectorError):
        read_archive(data)


def image_fixture(image_format="PNG"):
    import hashlib

    from openpyxl.drawing.image import Image
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from PIL import Image as PILImage

    source = BytesIO()
    PILImage.new("RGB", (2, 3), "red").save(source, format=image_format)
    media = source.getvalue()
    book = Workbook()
    book.active.title = "Report"
    book.calculation = None
    picture = Image(BytesIO(media))
    picture.anchor = OneCellAnchor(
        _from=AnchorMarker(col=0, row=5), ext=XDRPositiveSize2D(19050, 28575)
    )
    book.active.add_image(picture)
    out = BytesIO()
    book.save(out)
    expected = {
        "schema": "otc.xlsx-physical/1.0",
        "profile": "literal-artifact/1.0",
        "sheets": [
            {
                "name": "Report",
                "cells": {},
                "config": {},
                "merges": [],
                "images": [
                    {
                        "sha256": "sha256:" + hashlib.sha256(media).hexdigest(),
                        "mime_type": "image/png" if image_format == "PNG" else "image/jpeg",
                        "anchor": "A6",
                        "width_emu": 19050,
                        "height_emu": 28575,
                    }
                ],
            }
        ],
    }
    return out.getvalue(), expected


def test_independent_image_bytes_anchor_and_limits():
    data, expected = image_fixture()
    assert verify_snapshot(data, expected)["images"] == 1
    for limit in ("image_bytes", "total_image_bytes", "image_pixels"):
        with pytest.raises(ConnectorError):
            verify_snapshot(data, expected, {limit: 1})
    for old, new in ((b"<row>5</row>", b"<row>6</row>"), (b'cx="19050"', b'cx="19051"')):
        changed = corrupt(
            data, "xl/drawings/drawing1.xml", lambda b, old=old, new=new: b.replace(old, new)
        )
        with pytest.raises(ConnectorError):
            verify_snapshot(changed, expected)
    expected["sheets"][0]["images"][0]["sha256"] = "sha256:" + "0" * 64
    with pytest.raises(ConnectorError):
        verify_snapshot(data, expected)


@pytest.mark.parametrize(
    "replacement",
    [
        lambda b: b.replace(b'Id="rId2"', b'Id="rId1"'),
        lambda b: b.replace(
            b'Target="/xl/worksheets/sheet1.xml"',
            b'TargetMode="External" Target="https://example.com/sheet.xml"',
        ),
        lambda b: b.replace(
            b'Target="/xl/worksheets/sheet1.xml"', b'Target="/xl/worksheets/sheet9.xml"'
        ),
    ],
)
def test_relationship_id_external_and_missing_target(replacement):
    data, expected = fixture()
    with pytest.raises(ConnectorError):
        verify_snapshot(corrupt(data, "xl/_rels/workbook.xml.rels", replacement), expected)


def test_shared_string_reference_and_hidden_merge_value():
    data, expected = fixture()
    changed = corrupt(
        data,
        "xl/worksheets/sheet1.xml",
        lambda b: b.replace(b't="inlineStr"', b't="s"').replace(
            b"<is><t>=literal</t></is>", b"<v>999</v>"
        ),
    )
    with pytest.raises(ConnectorError):
        verify_snapshot(changed, expected)
    changed = corrupt(
        data,
        "xl/worksheets/sheet1.xml",
        lambda b: b.replace(b'r="A1"', b'r="B1"').replace(
            b"</sheetData>",
            b'</sheetData><mergeCells count="1"><mergeCell ref="A1:B1"/></mergeCells>',
        ),
    )
    cell = expected["sheets"][0]["cells"].pop("A1")
    expected["sheets"][0]["cells"]["B1"] = cell
    expected["sheets"][0]["merges"] = ["A1:B1"]
    with pytest.raises(ConnectorError) as error:
        verify_snapshot(changed, expected)
    assert error.value.safe_details["reason"] == "artifact.hidden_merged_value"


def test_cells_limit_exact_boundary_and_invalid_manifest():
    data, expected = fixture()
    assert verify_snapshot(data, expected, {"cells": 1})["cells"] == 1
    expected["sheets"][0]["cells"]["A1"]["undeclared"] = True
    with pytest.raises(ConnectorError):
        verify_snapshot(data, expected)


def test_frozen_manifest_and_quoted_print_area():
    from types import MappingProxyType

    from open_table_connector.local_files.spreadsheet_verify import _sheet_config

    data, expected = fixture()

    def freeze(value):
        if isinstance(value, dict):
            return MappingProxyType({k: freeze(v) for k, v in value.items()})
        if isinstance(value, list):
            return tuple(freeze(v) for v in value)
        return value

    assert verify_snapshot(data, freeze(expected))["cells"] == 1
    book = Workbook()
    ws = book.active
    ws.title = "O'Brien, Report"
    ws.print_area = ["A1:B3", "D1:E2"]
    assert _sheet_config(ws)["print_area"] == ["A1:B3", "D1:E2"]


def test_style_only_intent_is_verified():
    data, expected = fixture()
    expected["sheets"][0]["styles"] = {"B2": {"format": "General", "style": {"bold": True}}}
    with pytest.raises(ConnectorError):
        verify_snapshot(data, expected)
    expected["sheets"][0]["styles"]["B2"]["style"]["bold"] = False
    assert verify_snapshot(data, expected)["cells"] == 1


@pytest.mark.parametrize(
    "part, replacement",
    [
        (
            "docProps/core.xml",
            lambda b: b.replace(b"</cp:coreProperties>", b"<cp:unknown/></cp:coreProperties>"),
        ),
        ("xl/workbook.xml", lambda b: b.replace(b"</workbook>", b"<unknown/></workbook>")),
        (
            "xl/worksheets/sheet1.xml",
            lambda b: b.replace(b'<c r="A1"', b'<c undeclared="x" r="A1"'),
        ),
        ("xl/worksheets/sheet1.xml", lambda b: b.replace(b'<row r="1"', b'<row r="2"')),
        ("xl/workbook.xml", lambda b: b.replace(b'name="Report"', b'unknown="x" name="Report"')),
    ],
)
def test_closed_raw_metadata_and_cell_schema(part, replacement):
    data, expected = fixture()
    with pytest.raises(ConnectorError):
        verify_snapshot(corrupt(data, part, replacement), expected)


def test_jpeg_bytes_are_not_transcoded():
    data, expected = image_fixture("JPEG")
    assert verify_snapshot(data, expected)["images"] == 1
    with ZipFile(BytesIO(data)) as archive:
        media = archive.read("xl/media/image1.jpeg")
    changed = corrupt(data, "xl/media/image1.jpeg", lambda b: b + b"trailing bytes")
    with pytest.raises(ConnectorError):
        verify_snapshot(changed, expected)
    assert media.startswith(b"\xff\xd8")


def test_all_archive_and_text_limits_exact_boundary():
    data, expected = fixture()
    with ZipFile(BytesIO(data)) as archive:
        sizes = [len(archive.read(info)) for info in archive.infolist()]
    limits = {
        "archive_bytes": len(data),
        "zip_members": len(sizes),
        "member_bytes": max(sizes),
        "decompressed_bytes": sum(sizes),
        "text_bytes": len(b"=literal"),
    }
    for name, bound in limits.items():
        assert verify_snapshot(data, expected, {name: bound})["cells"] == 1
        with pytest.raises(ConnectorError) as error:
            verify_snapshot(data, expected, {name: bound - 1})
        assert error.value.code.value == "resource_limit_exceeded"
        assert error.value.safe_details["limit"] == name


def test_image_limit_exact_boundaries():
    data, expected = image_fixture()
    with ZipFile(BytesIO(data)) as archive:
        length = len(archive.read("xl/media/image1.png"))
    for name, bound in {
        "image_bytes": length,
        "total_image_bytes": length,
        "image_pixels": 6,
    }.items():
        assert verify_snapshot(data, expected, {name: bound})["images"] == 1
        with pytest.raises(ConnectorError) as error:
            verify_snapshot(data, expected, {name: bound - 1})
        assert error.value.safe_details["limit"] == name


def test_count_limit_boundaries():
    from copy import deepcopy

    from openpyxl import load_workbook

    data, expected = fixture()
    book = load_workbook(BytesIO(data))
    book.calculation = None
    book.active["B1"] = "second"
    book.active["B1"].number_format = "@"
    book.create_sheet("Second")
    expected["sheets"][0]["cells"]["B1"] = {
        **deepcopy(expected["sheets"][0]["cells"]["A1"]),
        "value": "second",
    }
    expected["sheets"].append(
        {"name": "Second", "cells": {}, "config": {}, "merges": [], "images": []}
    )
    out = BytesIO()
    book.save(out)
    data = out.getvalue()
    assert verify_snapshot(data, expected, {"sheets": 2, "cells": 2})["cells"] == 2
    for name in ("sheets", "cells"):
        with pytest.raises(ConnectorError) as error:
            verify_snapshot(data, expected, {name: 1})
        assert error.value.safe_details["limit"] == name


def test_forged_central_directory_size_is_rejected():
    import struct

    from open_table_connector.local_files.spreadsheet_verify import read_archive

    data, expected = fixture()
    damaged = bytearray(data)
    offset = damaged.index(b"PK\x01\x02")
    # Under-report uncompressed size while retaining compressed bytes/CRC.
    struct.pack_into("<I", damaged, offset + 24, 1)
    for verify in (
        lambda: verify_snapshot(bytes(damaged), expected),
        lambda: read_archive(bytes(damaged)),
    ):
        with pytest.raises(ConnectorError):
            verify()


def test_strict_chart_is_unsupported():
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, Reference

    data, expected = fixture()
    book = load_workbook(BytesIO(data))
    book.calculation = None
    chart = BarChart()
    chart.add_data(Reference(book.active, min_col=1, min_row=1, max_row=1))
    book.active.add_chart(chart, "D1")
    out = BytesIO()
    book.save(out)
    with pytest.raises(ConnectorError) as error:
        verify_snapshot(out.getvalue(), expected)
    assert error.value.safe_details["reason"] == "artifact.unsupported_part"


def test_image_count_limit_boundary():
    from copy import deepcopy

    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image

    data, expected = image_fixture()
    with ZipFile(BytesIO(data)) as archive:
        content = archive.read("xl/media/image1.png")
    book = load_workbook(BytesIO(data))
    book.calculation = None
    picture = Image(BytesIO(content))
    book.active.add_image(picture, "A6")
    expected["sheets"][0]["images"].append(deepcopy(expected["sheets"][0]["images"][0]))
    out = BytesIO()
    book.save(out)
    assert verify_snapshot(out.getvalue(), expected, {"images": 2})["images"] == 2
    with pytest.raises(ConnectorError) as error:
        verify_snapshot(out.getvalue(), expected, {"images": 1})
    assert error.value.safe_details["limit"] == "images"


def test_general_limits_precede_openpyxl_decode():
    from open_table_connector.local_files.spreadsheet_workbook import _preservation
    from open_table_connector.spreadsheets import ArtifactLimits

    def read_archive(data, limits):
        return _preservation(data, ArtifactLimits(**limits))

    data, _ = fixture()
    with pytest.raises(ConnectorError):
        read_archive(data, {"text_bytes": 1})
    image_data, _ = image_fixture()
    with pytest.raises(ConnectorError) as error:
        read_archive(image_data, {"image_pixels": 1})
    assert error.value.safe_details["reason"] == "image_pixels"


def test_forged_size_with_matching_prefix_crc_cannot_hide_expansion():
    import struct
    import zlib

    from open_table_connector.local_files.spreadsheet_verify import read_archive

    data, expected = fixture()
    with ZipFile(BytesIO(data)) as archive:
        original = archive.read("docProps/core.xml")
    changed = corrupt(data, "docProps/core.xml", lambda b: b + b" " * 200000)
    damaged = bytearray(changed)
    with ZipFile(BytesIO(changed)) as archive:
        local_offset = archive.getinfo("docProps/core.xml").header_offset
    for offset in range(len(damaged) - 46):
        if damaged[offset : offset + 4] == b"PK\x01\x02":
            length = struct.unpack_from("<H", damaged, offset + 28)[0]
            if damaged[offset + 46 : offset + 46 + length] == b"docProps/core.xml":
                struct.pack_into("<I", damaged, offset + 16, zlib.crc32(original))
                struct.pack_into("<I", damaged, offset + 24, len(original))
    struct.pack_into("<I", damaged, local_offset + 14, zlib.crc32(original))
    struct.pack_into("<I", damaged, local_offset + 22, len(original))
    # Python's ordinary ZIP reader accepts this forged prefix and discards expansion.
    with ZipFile(BytesIO(damaged)) as archive:
        assert archive.read("docProps/core.xml") == original
    for verify in (
        lambda: verify_snapshot(bytes(damaged), expected, {"member_bytes": 20000}),
        lambda: read_archive(bytes(damaged), {"member_bytes": 20000}),
    ):
        with pytest.raises(ConnectorError) as error:
            verify()
        assert error.value.code.value == "resource_limit_exceeded"
